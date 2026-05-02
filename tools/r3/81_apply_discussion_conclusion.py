"""
R3 final wording pass 2 — Discussion / Conclusion / Limitations.

Applies the user's chốt blocks for Section 5.1, 5.2, 6, and the
Section 7 limitation 2 paragraph, plus three short sweep edits to
keep tone consistent with the new Section 4.2 framing.

Changes (in document order):
  para  74  Section 4.3 closing — soften "level-shift effect"
  para  91  Figure 3 explanatory — drop "extreme-case Singapore"
  para  95  Section 4.5 closing — drop "digitally mature setting"
  paras 103-105  Section 5.1 — full rewrite (3 paragraphs)
  paras 107-109  Section 5.2 — full rewrite (3 paragraphs)
  paras 111-113  Section 6 Conclusion — full rewrite (3 paragraphs)
  para  116      Section 7 limitation 2 — full replacement
"""
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers/p3-singapore/Manuscript_Blinded_MIR_2_revised.docx"


# ============================================================
# Section 5.1 (3 paragraphs)
# ============================================================
S51_P1 = (
    "Three theoretical implications follow from these findings. First, "
    "the TCI result supports a capability-depth interpretation of firm "
    "performance in internationalizing firms. Technological capability "
    "is positively associated with labour productivity in a way that is "
    "more stable than any moderation pattern identified in the present "
    "design, which is consistent with the view that internal capability "
    "depth improves the productivity base from which firms engage "
    "international markets. This interpretation fits the absorptive-"
    "capacity and technological-capability traditions by emphasizing "
    "learning, innovation, and technology absorption as firm-internal "
    "sources of productivity heterogeneity rather than as clearly "
    "identified shifters of I–P curvature."
)
S51_P2 = (
    "Second, the evidence qualifies rather than overturns the "
    "conventional nonlinear I–P literature. In the Singapore sample, "
    "the fitted quadratic displays mild curvature, but the right-side "
    "decline is not formally identified within the export-intensity "
    "range occupied by most firms, and the implied turning point lies "
    "in a sparsely populated upper tail. The appropriate theoretical "
    "reading is therefore not that this study establishes a general "
    "boundary condition for digitally mature economies, but that it "
    "provides within-context evidence from Singapore showing how the "
    "conventional I–P logic appears when most firms remain clustered "
    "at very low export intensity and the upper tail is thin. In that "
    "sense, the study sharpens interpretation of the nonlinear "
    "literature without claiming that a single-country cross-section "
    "can separate a digital-frontier mechanism from Singapore-specific "
    "institutional features."
)
S51_P3 = (
    "Third, the DAI result suggests that foundational digital adoption "
    "is better understood as a conditional scaling resource than as a "
    "uniform firm-level productivity premium. Across much of the "
    "observed export-intensity distribution, the DAI association is "
    "weak or statistically indistinct, but it becomes more positive in "
    "the high-export tail, where firms face denser cross-border "
    "coordination and transaction demands. This pattern is consistent "
    "with the idea that Tier 1–2 digital adoption matters most when "
    "firms have sufficient international throughput to use digital "
    "interfaces and transaction-enabling systems intensively, while "
    "also underscoring that the evidence is concentrated in a thin-"
    "support region and should therefore be interpreted cautiously. "
    "The contribution here is thus not to establish a universal "
    "digitalization mechanism, but to clarify that foundational "
    "digital adoption and technological capability are analytically "
    "distinct and empirically associated with performance in different "
    "ways within this setting."
)

# ============================================================
# Section 5.2 (3 paragraphs)
# ============================================================
S52_P1 = (
    "For firms in Singapore and comparable high-digital-infrastructure "
    "contexts, the findings suggest that investments in foundational "
    "digital adoption are most strongly associated with productivity "
    "advantages when firms already operate at relatively high export "
    "intensity. In such cases, digital interfaces, payment systems, "
    "and related transaction-enabling tools appear to function as "
    "scaling mechanisms that help firms coordinate a larger volume of "
    "cross-border activity more efficiently. For firms already deeply "
    "engaged in export markets, the managerial implication is "
    "therefore strategic rather than symbolic: foundational digital "
    "systems appear to matter most when coordination demands across "
    "customers, suppliers, and foreign-market transactions become "
    "dense enough for those systems to be used intensively."
)
S52_P2 = (
    "For firms concentrated in domestic markets or at low export "
    "intensity, digital adoption may still be worthwhile, but the "
    "productivity differentials appear smaller and less distinctive in "
    "the present data. In this setting, basic digital functionality is "
    "already relatively widespread, so adopting such tools may not "
    "generate a large standalone labour-productivity premium for all "
    "firms equally. For firms in the intermediate export-intensity "
    "range, the results suggest caution in expecting immediate "
    "productivity gains from digital adoption alone; these investments "
    "are likely to be more valuable when aligned with broader "
    "objectives such as customer integration, supplier coordination, "
    "export readiness, and future scaling."
)
S52_P3 = (
    "More broadly, the findings imply that managers should avoid "
    "treating technological capability and digital adoption as "
    "interchangeable investment categories. Investments in "
    "technological capability appear to support a broader productivity "
    "base, whereas investments in foundational digital adoption become "
    "more relevant when firms face the coordination intensity "
    "associated with deeper internationalization. This suggests that "
    "capability-building and digitalization strategies should be "
    "sequenced and matched to the firm's actual stage of export "
    "expansion rather than pursued as a single undifferentiated "
    "digital-capability agenda."
)

# ============================================================
# Section 6 Conclusion (3 paragraphs)
# ============================================================
S6_P1 = (
    "This study revisits the internationalization–performance "
    "relationship by distinguishing technological capability from "
    "foundational digital adoption and examining how each is "
    "associated with labour productivity among firms in Singapore. "
    "Using World Bank Enterprise Survey microdata for Singapore 2023, "
    "the analysis shows that technological capability is positively "
    "associated with productivity, while digital adoption exhibits a "
    "more conditional association that becomes more positive only at "
    "higher levels of export intensity. The findings therefore "
    "support a distinction between firm-internal capability depth and "
    "foundational digitally enabled transaction capacity rather than "
    "treating both domains as interchangeable aspects of a single "
    "digital-capability construct."
)
S6_P2 = (
    "The study also qualifies the interpretation of the nonlinear "
    "I–P literature in this setting. Within the observed export-"
    "intensity range, the baseline pattern is better characterized as "
    "predominantly positive with mild quadratic curvature than as a "
    "formally identified inverted-U, because the implied turning point "
    "lies in a sparsely populated upper tail and is imprecisely "
    "located. Read in this way, the evidence does not overturn the "
    "established I–P literature, nor does it establish a general "
    "boundary condition for digitally mature economies; instead, it "
    "provides within-context evidence from Singapore on how the "
    "conventional nonlinear logic appears in a highly digitalized "
    "institutional environment where most firms remain concentrated "
    "at low export intensity."
)
S6_P3 = (
    "More broadly, the study contributes to international business "
    "research by sharpening construct interpretation and by showing "
    "that technological capability and foundational digital adoption "
    "are associated with performance through different empirical "
    "patterns. Technological capability shows the more stable positive "
    "association with productivity, whereas foundational digital "
    "adoption is better interpreted as a conditional scaling resource "
    "whose relevance becomes more visible only where cross-border "
    "coordination demands are relatively intense. These conclusions "
    "are deliberately bounded to the present design and setting, and "
    "they point toward the need for comparative and longitudinal "
    "research before broader claims about digital-frontier mechanisms "
    "can be sustained."
)

# ============================================================
# Section 7 limitation 2 (full replacement)
# ============================================================
S7_LIM2 = (
    "Second, the sample contains a thin right tail of high-intensity "
    "exporters. Most firms report zero exports, the 75th percentile of "
    "FSTS remains at zero, and only a small fraction of firms occupy "
    "the high-export range in which the fitted turning point and the "
    "strongest DAI moderation signals appear. Two consequences follow. "
    "First, the implied turning point in the quadratic specification "
    "should be interpreted as an indicative descriptive feature rather "
    "than as a structurally identified inflection, because it lies in "
    "a sparsely populated region and its bootstrap confidence interval "
    "is wide ([53%, 253%]) even though an inverted-U shape is "
    "frequently recovered in resamples (96.3% of replications). "
    "Second, the DAI moderation pattern is more appropriately treated "
    "as suggestive evidence concentrated in the thin upper tail: it "
    "remains visible across leave-one-out, trimmed-tail, and "
    "indicator-sensitivity diagnostics, but precision is necessarily "
    "lower when attention shifts to the small exporter subsample "
    "(N = 84) and the sparsely populated high-intensity range. The "
    "appropriate caution therefore rests on the limited empirical "
    "support in the upper tail and the resulting imprecision of tail-"
    "based inference, not on differences in adjusted R² across "
    "samples of very different size and composition."
)

# ============================================================
# Sweep edits (single phrase changes)
# ============================================================
PARA74_OLD = (
    "the manuscript therefore should interpret TCI primarily as a "
    "level-shift effect on the productivity intercept rather than as "
    "a curvature-shaping moderator"
)
PARA74_OLD_ALT = (
    "interpret TCI primarily as a level-shift effect on the "
    "productivity intercept rather than as a curvature-shaping moderator"
)
PARA74_NEW = (
    "the evidence is therefore more consistent with an intercept-"
    "dominant reading than with a clearly identified moderation channel"
)

PARA91_OLD = "in this extreme-case Singapore setting"
PARA91_NEW = "in this Singapore sample"

PARA95_OLD = (
    "its association with productivity becomes more positive as "
    "export intensity rises in a digitally mature setting"
)
PARA95_NEW = (
    "its association with productivity becomes more positive as "
    "export intensity rises within the Singapore sample, with the "
    "clearest signal concentrated in the high-export tail"
)


def replace_para(p, new_text: str):
    if not p.runs:
        p.text = new_text
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def sweep(p, old, new):
    if old in p.text:
        new_t = p.text.replace(old, new)
        replace_para(p, new_t)
        return True
    return False


def main():
    doc = Document(DOC)

    print("[1] Replacing Section 5.1 (paras 103-105)")
    replace_para(doc.paragraphs[103], S51_P1)
    replace_para(doc.paragraphs[104], S51_P2)
    replace_para(doc.paragraphs[105], S51_P3)

    print("[2] Replacing Section 5.2 (paras 107-109)")
    replace_para(doc.paragraphs[107], S52_P1)
    replace_para(doc.paragraphs[108], S52_P2)
    replace_para(doc.paragraphs[109], S52_P3)

    print("[3] Replacing Section 6 Conclusion (paras 111-113)")
    replace_para(doc.paragraphs[111], S6_P1)
    replace_para(doc.paragraphs[112], S6_P2)
    replace_para(doc.paragraphs[113], S6_P3)

    print("[4] Replacing Section 7 limitation 2 (para 116)")
    replace_para(doc.paragraphs[116], S7_LIM2)

    print("[5] Sweep para 74 (Section 4.3 closing soften)")
    if not sweep(doc.paragraphs[74], PARA74_OLD, PARA74_NEW):
        # try alt phrasing
        if sweep(doc.paragraphs[74], PARA74_OLD_ALT, PARA74_NEW):
            print("    matched ALT phrasing")
        else:
            print("    (no exact match — printing for manual review)")
            print(f"    actual: {doc.paragraphs[74].text[:300]}")

    print("[6] Sweep para 91 (Figure 3 explanatory)")
    if not sweep(doc.paragraphs[91], PARA91_OLD, PARA91_NEW):
        print(f"    (no match — actual: {doc.paragraphs[91].text[:200]})")

    print("[7] Sweep para 95 (Section 4.5 closing)")
    if not sweep(doc.paragraphs[95], PARA95_OLD, PARA95_NEW):
        print(f"    (no exact match — actual: {doc.paragraphs[95].text[:300]})")

    doc.save(DOC)
    print(f"\nSaved: {DOC}")


if __name__ == "__main__":
    main()
