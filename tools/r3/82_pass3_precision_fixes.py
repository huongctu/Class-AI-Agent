"""
R3 Pass 3 — Six precision fixes per co-author review:

  1. Keywords: drop "boundary conditions" → swap with "export intensity;
     labour productivity".
  2. H2 statement (para 35): demote moderation claim to research-question
     style; remove duplication with H1.
  3. Section 2.4 (para 43) + Figure 1 caption (para 45): reflect that DAI
     is the contingency variable on the I-P path, while TCI enters
     mainly as a direct-effect construct with supplementary moderation.
  4. Section 3.3 (para 61): replace "operates primarily as a moderator
     or as an intercept-shifting capability variable" with a softer
     supplementary-test framing.
  5. Section 4.3 (para 74): repair broken sentence
     "The manuscript therefore should the evidence is therefore more
     consistent..." (left-over from previous sweep that only matched
     part of the original sentence).
  6. Section 4.4 numerical consistency:
     6a. Para 82: text reports R² = 0.213, adj R² = 0.199 but Table 2
         M8 reports 0.211 / 0.196 — fix text to match table.
     6b. Para 87: text says marginal effect of DAI is indistinguishable
         from zero "across domestic, low-export, and much of the middle
         export range" but Table 4 shows FSTS=0 has p=.045 (sig).
         Reword to acknowledge the small positive domestic effect
         before noting null mid-range effects.
"""
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers/p3-singapore/Manuscript_Blinded_MIR_2_revised.docx"


# =================================================================
# Replacement strings
# =================================================================

NEW_KEYWORDS = (
    "Keywords: internationalization–performance relationship; "
    "digital adoption; technological capability; export intensity; "
    "labour productivity; Singapore"
)

# H2 statement (para 35) — demote to RQ; H1 already covers TCI direct effect
NEW_H2 = (
    "Whether technological capability (TCI) also moderates the "
    "internationalization–performance relationship is treated as an "
    "open empirical question and assessed in a supplementary "
    "specification rather than as a hypothesized moderation channel."
)

# Section 2.4 content (para 43) — TCI direct, DAI contingency
NEW_S24 = (
    "Figure 1 summarizes the conceptual model. The model treats TCI "
    "and DAI as analytically distinct firm-level constructs that are "
    "associated with firm performance through different channels. TCI "
    "enters the model primarily as a direct-effect construct, "
    "reflecting firm-internal capability depth grounded in learning, "
    "innovation, and technology absorption; whether TCI also "
    "moderates the FSTS–performance relationship is examined in a "
    "supplementary test rather than presumed ex ante. DAI, by "
    "contrast, is the contingency variable on the I–P path: its "
    "productivity relevance is theorized to vary across levels of "
    "export intensity rather than to operate as a uniform direct "
    "premium, because the observed indicators capture foundational "
    "digital interfaces and transaction-enabling mechanisms whose "
    "relevance depends on how intensively firms engage in cross-"
    "border operations. In this sense, H3 and H4 are jointly "
    "informative: H3 concerns the non-uniformity of the DAI–"
    "productivity association, whereas H4 specifies the direction of "
    "that contingency across export intensity. The model is therefore "
    "positioned as a within-context framework for interpreting firm-"
    "level evidence from Singapore rather than as a design that can, "
    "on its own, establish a general boundary condition for digitally "
    "mature economies."
)

# Figure 1 caption (para 45)
NEW_F1_CAPTION = (
    "Figure 1. Conceptual model. Internationalization (FSTS, FSTS²) "
    "is the independent variable; firm performance, measured as "
    "ln(labour productivity), is the dependent variable. Digital "
    "adoption (DAI) is the contingency variable on the I–P path "
    "(H4); technological capability (TCI) enters mainly as a direct-"
    "effect construct (H1), with any moderation by TCI assessed in a "
    "supplementary test rather than hypothesized ex ante. The diagram "
    "also lists the direct effect of DAI on the dependent variable "
    "(H3), which is estimated in the empirical model but not drawn as "
    "a crossing arrow for visual clarity. Firm size, age, foreign "
    "ownership, and broad-sector fixed effects enter as controls. "
    "Solid arrows denote direct associations; dashed arrows denote "
    "moderation associations. The setting is Singapore as a within-"
    "context, analytically informative case of a digitally advanced "
    "economy (WBES 2023, N = 623 / 617)."
)

# Section 3.3 (para 61) — soften "operates primarily as a moderator..."
S33_OLD_PHRASE = (
    "in order to assess whether technological capability operates "
    "primarily as a moderator or as an intercept-shifting capability "
    "variable"
)
S33_NEW_PHRASE = (
    "in order to assess whether any TCI moderation is empirically "
    "detectable beyond its more stable direct association with "
    "productivity"
)

# Section 4.3 (para 74) — repair broken sentence
S43_OLD_FRAGMENT = (
    "The manuscript therefore should the evidence is therefore more "
    "consistent with an intercept-dominant reading than with a "
    "clearly identified moderation channel of the internationalization–"
    "performance relationship."
)
S43_NEW_SENTENCE = (
    "Taken together, the evidence is more consistent with an "
    "intercept-dominant reading of the TCI association than with a "
    "clearly identified moderation channel in the internationalization–"
    "performance relationship."
)

# Section 4.4 — fix R² numbers (para 82)
S44_OLD_RSQ = "with R² = 0.213 and adjusted R² = 0.199"
S44_NEW_RSQ = "with R² = 0.211 and adjusted R² = 0.196"

# Section 4.4 — fix domestic-effect contradiction (para 87)
S44_OLD_DOMESTIC = (
    "Across domestic, low-export, and much of the middle export "
    "range, the marginal effect of DAI is small and statistically "
    "indistinguishable from zero, which indicates that basic digital "
    "adoption does not generate a large universal labour-productivity "
    "premium across the full sample."
)
S44_NEW_DOMESTIC = (
    "At FSTS = 0, the marginal effect of DAI is small but positive "
    "and marginally significant (+0.080, p = .045), indicating only "
    "a modest baseline association among purely domestic firms. "
    "Across the low-export and middle export range, however, the "
    "marginal effect is statistically indistinguishable from zero, "
    "which indicates that basic digital adoption does not generate a "
    "large universal labour-productivity premium across the full "
    "sample."
)


def replace_para(p, new_text: str):
    if not p.runs:
        p.text = new_text
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def sweep(p, old, new, label):
    if old in p.text:
        new_t = p.text.replace(old, new)
        replace_para(p, new_t)
        print(f"      ✓ {label}")
        return True
    print(f"      ✗ NO MATCH for {label}")
    print(f"        Looking for: '{old[:80]}...'")
    print(f"        In: '{p.text[:200]}...'")
    return False


def main():
    doc = Document(DOC)

    print("[1] Update Keywords (para 5)")
    replace_para(doc.paragraphs[5], NEW_KEYWORDS)
    print(f"      ✓ Keywords now: '{NEW_KEYWORDS[10:60]}...'")

    print("\n[2] Demote H2 statement (para 35)")
    replace_para(doc.paragraphs[35], NEW_H2)
    print(f"      ✓ H2 reframed as research question")

    print("\n[3a] Section 2.4 content (para 43)")
    replace_para(doc.paragraphs[43], NEW_S24)
    print(f"      ✓ TCI = direct, DAI = contingency on I-P path")

    print("\n[3b] Figure 1 caption (para 45)")
    replace_para(doc.paragraphs[45], NEW_F1_CAPTION)
    print(f"      ✓ Caption updated to match new framework")

    # Also update duplicate caption at end of doc (was para 156 originally)
    for i in range(150, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().startswith("Figure 1. Conceptual model"):
            replace_para(doc.paragraphs[i], NEW_F1_CAPTION)
            print(f"      ✓ Duplicate Figure 1 caption at para {i} also updated")
            break

    print("\n[4] Section 3.3 supplementary phrasing (para 61)")
    sweep(doc.paragraphs[61], S33_OLD_PHRASE, S33_NEW_PHRASE,
          "Section 3.3 supplementary spec phrase softened")

    print("\n[5] Section 4.3 broken sentence repair (para 74)")
    sweep(doc.paragraphs[74], S43_OLD_FRAGMENT, S43_NEW_SENTENCE,
          "Section 4.3 grammar bug fixed")

    print("\n[6a] Section 4.4 R² numerical fix (para 82)")
    sweep(doc.paragraphs[82], S44_OLD_RSQ, S44_NEW_RSQ,
          "M8 R² aligned with Table 2 (0.211 / 0.196)")

    print("\n[6b] Section 4.4 domestic-effect rewording (para 87)")
    sweep(doc.paragraphs[87], S44_OLD_DOMESTIC, S44_NEW_DOMESTIC,
          "Domestic-effect text now consistent with Table 4")

    doc.save(DOC)
    print(f"\nSaved: {DOC}")


if __name__ == "__main__":
    main()
