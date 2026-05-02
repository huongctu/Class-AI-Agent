"""
R3 Pass 4 — final cleanup before resubmission.

  1. Replace Section 7 (paras 115, 116, 117) with the shorter, cleaner
     version the user provided in the editor-simulation review. Drops
     the "extreme-case, within-context evidence from Singapore"
     phrasing in limitation 1, removes the long Heckman/equivalence
     testing tangents, and re-anchors limitation 3 in the Tier 1-2
     measurement boundary + future-research priorities.

  2. Normalize "digitally mature" → "digitally advanced" globally,
     except where the phrase appears inside an explicit negation of
     the boundary-condition claim (which is rhetorically essential).

  3. Para 10 (Section 1.1): soft sweep — "In digitally mature
     settings" → "In digitally advanced settings".

  4. Para 8 (Section 1.1): "those of digitally mature economies" →
     "those of digitally advanced economies".
"""
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers/p3-singapore/Manuscript_Blinded_MIR_2_revised.docx"


# =================================================================
# Section 7 — three replacement paragraphs (shorter + cleaner)
# =================================================================
S7_LIM1 = (
    "Three limitations warrant caution in interpreting the findings. "
    "First, the analysis is based on a single-country cross-section, "
    "so the identification scope is associational rather than causal. "
    "The observed relationships among technological capability, "
    "digital adoption, export intensity, and labour productivity may "
    "reflect reverse causation, omitted variables, or productivity-"
    "driven selection into exporting and digital adoption. The "
    "appropriate interpretation is therefore that the study documents "
    "how these variables co-vary in Singapore rather than establishing "
    "a causal mechanism."
)

S7_LIM2 = (
    "Second, the sample contains a thin right tail of high-intensity "
    "exporters. Most firms report zero exports, the 75th percentile of "
    "FSTS remains at zero, and only a small fraction of firms occupy "
    "the high-export range in which the fitted turning point and the "
    "strongest DAI moderation signals appear. The implied turning "
    "point in the quadratic specification should therefore be "
    "interpreted as a descriptive feature rather than as a structurally "
    "identified inflection, because it is estimated from a sparsely "
    "populated region and its 95% bootstrap confidence interval is "
    "wide ([53%, 253%]) even though an inverted-U shape is recovered "
    "in 96.3% of resamples. The same caution applies to the DAI "
    "moderation pattern: it remains visible across leave-one-out, "
    "trimmed-tail, and indicator-sensitivity diagnostics, but "
    "precision is necessarily lower in the small exporter subsample "
    "(N = 84) and the sparsely populated high-intensity range. The "
    "appropriate caution therefore rests on the limited empirical "
    "support in the upper tail and the resulting imprecision of tail-"
    "based inference."
)

S7_LIM3 = (
    "Third, the estimated digital-adoption effect should be "
    "interpreted within the measurement boundary of the available "
    "indicators. The DAI construct captures Tier 1–2 digital adoption "
    "— digital presence and electronic-transaction usage — rather "
    "than deeper digitally integrated organizational capability or "
    "digital dynamic capability; indicator-sensitivity analysis shows "
    "that the moderation pattern depends on the foundational digital-"
    "infrastructure indicators and weakens when those specific items "
    "are dropped. Future research should therefore examine whether "
    "the same conditional pattern appears in other digitally advanced "
    "economies and under richer measures that capture process "
    "integration, organizational digitalization, and higher-order "
    "digital capabilities more directly. Panel data, successive "
    "enterprise-survey waves, linked administrative records, and "
    "stronger quasi-experimental designs — including instrumental-"
    "variable or difference-in-differences specifications keyed to "
    "digital-infrastructure rollouts or trade-policy shocks — would "
    "also help clarify whether the observed associations reflect "
    "scaling complementarities, selection effects, or other "
    "mechanisms."
)


# =================================================================
# Targeted "digitally mature" sweeps
# (Only positive contexts. Negation contexts disavowing the
#  boundary-condition claim are preserved verbatim.)
# =================================================================
SWEEPS = [
    # Para 8 — Section 1.1 (positive context)
    ("differ substantially from those of digitally mature economies",
     "differ substantially from those of digitally advanced economies"),
    # Para 10 — Section 1.1
    ("In digitally mature settings, firms may differ",
     "In digitally advanced settings, firms may differ"),
    # Para 115 — Section 7 limitation 1 (will be fully replaced anyway,
    # but listed for completeness)
    ("co-move with export intensity in a digitally mature economy",
     "co-move with export intensity in this digitally advanced setting"),
    # Para 117 — Section 7 limitation 3 (will be fully replaced anyway)
    ("in one digitally mature setting",
     "in one digitally advanced setting"),
    ("the same pattern appears in other digitally mature economies",
     "the same pattern appears in other digitally advanced economies"),
]


def replace_para(p, new_text: str):
    if not p.runs:
        p.text = new_text
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def main():
    doc = Document(DOC)

    print("[1] Replacing Section 7 limitation 1 (para 115)")
    replace_para(doc.paragraphs[115], S7_LIM1)
    print(f"      Old length → new length: 223 → {len(S7_LIM1.split())} words")

    print("\n[2] Replacing Section 7 limitation 2 (para 116)")
    replace_para(doc.paragraphs[116], S7_LIM2)
    print(f"      Old length → new length: 189 → {len(S7_LIM2.split())} words")

    print("\n[3] Replacing Section 7 limitation 3 (para 117)")
    replace_para(doc.paragraphs[117], S7_LIM3)
    print(f"      Old length → new length: 313 → {len(S7_LIM3.split())} words")

    print("\n[4] Sweep 'digitally mature' → 'digitally advanced' "
          "(positive contexts only)")
    n_changes = 0
    for old, new in SWEEPS:
        for p in doc.paragraphs:
            if old in p.text:
                replace_para(p, p.text.replace(old, new))
                n_changes += 1
                print(f"      ✓ '{old[:60]}...' → '{new[:60]}...'")
    print(f"\n      Total sweeps applied: {n_changes}")

    doc.save(DOC)
    print(f"\nSaved: {DOC}")

    # Word count after fix
    doc = Document(DOC)
    sec7_start = next(i for i, p in enumerate(doc.paragraphs)
                      if p.text.strip().startswith("7 Limitations"))
    sec7_end = next(i for i, p in enumerate(doc.paragraphs)
                    if p.text.strip().startswith("Data Availability"))
    sec7_words = sum(len(p.text.split())
                     for p in doc.paragraphs[sec7_start:sec7_end])
    total = sum(len(p.text.split()) for p in doc.paragraphs)
    print(f"\nSection 7 length after Pass 4: {sec7_words} words "
          f"(was 730 words)")
    print(f"Total manuscript: {total:,} words")


if __name__ == "__main__":
    main()
