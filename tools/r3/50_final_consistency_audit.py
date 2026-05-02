"""
R3 Phase 8 — Final consistency audit across the entire submission package.

Checks:
  1. Title is identical across manuscript, cover letter, title page.
  2. No "digital-frontier" mentions in any submission file.
  3. No stale Table 3 numbers (0.187, 2.972, 1.308, etc.) in text outside
     of Table 3 itself.
  4. Manuscript Table 3 matches canonical re-estimation
     (outputs/r3/audit/table3_audit.json).
  5. Abstract word count ≤ 250.
  6. References section intact; in-text citations follow MIR style
     "(Author Year)" with no comma.
  7. Hypothesis numbering is consistent (H1–H4).
  8. No leftover author info / blinding leaks.

Reports any issues found.
"""
from pathlib import Path
import json
import re
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers" / "p3-singapore" / "Manuscript_Blinded_MIR_2_revised.docx"
COVER = REPO / "papers" / "p3-singapore" / "Cover_Letter.docx"
TITLE_PAGE = REPO / "papers" / "p3-singapore" / "Title_Page.docx"
TABLE3_AUDIT = REPO / "outputs" / "r3" / "audit" / "table3_audit.json"


def all_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def main():
    print("=" * 72)
    print("R3 Phase 8 — Final consistency audit")
    print("=" * 72)
    issues = []
    passes = []

    doc = Document(DOC)
    cov = Document(COVER)
    tp = Document(TITLE_PAGE)
    text_doc = all_text(doc)
    text_cov = all_text(cov)
    text_tp = all_text(tp)

    # 1. Title consistency
    title_doc = doc.paragraphs[0].text.strip()
    title_in_cover = title_doc in text_cov
    # In title page, title is typically a Run with bold formatting
    title_in_tp = title_doc in text_tp
    if title_in_cover and title_in_tp:
        passes.append(f"✓ Title identical across all 3 files")
    else:
        issues.append(f"✗ Title mismatch — manuscript title not found in "
                      f"{'cover letter' if not title_in_cover else ''} "
                      f"{'title page' if not title_in_tp else ''}")
    print(f"\n[1] Title across files:")
    print(f"    Manuscript: {title_doc[:90]}...")

    # 2. No digital-frontier in any file
    for label, t in [("manuscript", text_doc), ("cover letter", text_cov),
                     ("title page", text_tp)]:
        n = len(re.findall(r"digital[-\s]?frontier", t, re.I))
        if n == 0:
            passes.append(f"✓ No 'digital-frontier' in {label}")
        else:
            issues.append(f"✗ {n} 'digital-frontier' mentions in {label}")
    print(f"\n[2] 'digital-frontier' sweep: see passes/issues below")

    # 3. Stale Table 3 numbers in text (not cells)
    text_outside_tables = "\n".join(
        p.text for p in doc.paragraphs)  # paragraphs already exclude tables
    stale_patterns = [
        (r"\b0\.187\*\*\*", "old Table 3 baseline TCI"),
        (r"\b2\.972\*", "old Table 3 baseline FSTS²×DAI"),
        (r"\b1\.308\b", "old Table 3 R5 FSTS²×DAI"),
        (r"\b0\.218\*\*\*", "old Table 3 R3 TCI"),
    ]
    for pat, desc in stale_patterns:
        if re.search(pat, text_outside_tables):
            issues.append(f"✗ Stale number found in text: {desc}")
        else:
            passes.append(f"✓ No stale '{desc}' in text")

    # 4. Manuscript Table 3 matches canonical (numeric content only;
    #    ignore "+" sign and "0.001" vs ".001" formatting differences).
    def normalize(s):
        return (s.lstrip("+").replace("(0.", "(.").replace(" 0.", " ."))

    canon = json.loads(TABLE3_AUDIT.read_text())
    table3 = doc.tables[3]
    canon_rows = canon["corrected_table"]
    manus_rows = []
    for ri in range(1, 7):
        cells = [table3.rows[ri].cells[j].text.strip() for j in range(6)]
        manus_rows.append(tuple(normalize(c) for c in cells))
    canon_tuples = [tuple(normalize(c) for c in r) for r in canon_rows]
    if manus_rows == canon_tuples:
        passes.append("✓ Manuscript Table 3 matches canonical re-estimation "
                      "(numeric content)")
    else:
        issues.append("✗ Manuscript Table 3 does NOT match canonical")
        for i, (m, c) in enumerate(zip(manus_rows, canon_tuples)):
            if m != c:
                issues.append(f"   Row {i+1} differs:")
                issues.append(f"     Manuscript: {m}")
                issues.append(f"     Canonical:  {c}")

    # 5. Abstract word count
    abs_paras = []
    started = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "Abstract":
            started = True
            continue
        if started:
            if t.startswith(("Keywords", "1 Introduction")):
                break
            if t:
                abs_paras.append(t)
    abs_words = len(" ".join(abs_paras).split())
    if abs_words <= 250:
        passes.append(f"✓ Abstract = {abs_words} words (≤ 250)")
    else:
        issues.append(f"✗ Abstract = {abs_words} words (over 250)")

    # 6. References + MIR citation style
    if "References" in text_doc:
        passes.append("✓ References section present")
    else:
        issues.append("✗ References section missing")

    apa_pattern = re.compile(
        r"\(([A-Z][A-Za-z'’\-]+(?:\set\sal\.)?), (\d{4})\)")
    apa_matches = apa_pattern.findall(text_doc)
    if apa_matches:
        issues.append(f"✗ {len(apa_matches)} APA-style citations remain "
                      f"(MIR uses no comma): e.g., {apa_matches[:3]}")
    else:
        passes.append("✓ All in-text citations use MIR style (Author Year)")

    # 7. Hypothesis numbering
    h_present = []
    for h in ("H1", "H2", "H3", "H4"):
        if re.search(rf"Hypothesis {h[1]} \({h}\)", text_doc) or \
           re.search(rf"\b{h}\b", text_doc):
            h_present.append(h)
    if len(h_present) == 4:
        passes.append(f"✓ All four hypotheses (H1–H4) present")
    else:
        issues.append(f"✗ Hypothesis coverage incomplete: only {h_present}")

    # 8. Blinding check
    leaks = []
    for pattern in [r"@gmail|@yahoo|@hotmail|@nus\.edu|@ntu\.edu",
                    r"University of [A-Z][a-z]+",
                    r"ORCID:\s?\d{4}",
                    r"Acknowledgments?:\s*[A-Z]"]:
        m = re.search(pattern, text_doc)
        if m:
            leaks.append(f"   '{m.group(0)}'")
    if leaks:
        issues.append("✗ Possible blinding leaks:")
        issues.extend(leaks)
    else:
        passes.append("✓ No author identifiers detected (blinded)")

    # Summary
    print("\n" + "=" * 72)
    print(f"PASSES ({len(passes)}):")
    print("=" * 72)
    for p in passes:
        print(f"  {p}")
    print("\n" + "=" * 72)
    print(f"ISSUES ({len(issues)}):")
    print("=" * 72)
    if not issues:
        print("  None — submission package is consistent.")
    else:
        for i in issues:
            print(f"  {i}")

    print("\n" + "=" * 72)
    print(f"VERDICT: {'READY FOR SUBMISSION' if not issues else 'FIXES NEEDED'}")
    print("=" * 72)


if __name__ == "__main__":
    main()
