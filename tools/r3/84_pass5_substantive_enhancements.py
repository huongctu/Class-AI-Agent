"""
R3 Pass 5 — Three substantive enhancements per NotebookLM editor review:

  A. Section 4.5 (para 94): Add item-swap falsification disclosure to make
     the "ultimate weapon" construct-validity test visible in the
     manuscript robustness section. Currently item-swap is documented in
     audit JSON + Master Memo but NOT in manuscript prose.

  B. Section 7 (para 117): Expand the future-research paragraph with
     concrete instrumental-variable suggestions (firm-to-port distance,
     customs-clearance days, regional broadband coverage, within-sector
     peer adoption), with Wolfolds & Siegel (2019) caveat about
     first-stage strength and exclusion restrictions.

  C. Section 5.1 implication 3 (para 105): Add transaction-cost-economics
     (Coase 1937; Williamson 1985) framing for the conditional DAI
     pattern — explicit theoretical mechanism for why digital adoption
     matters more at high export intensity.

All three additions extend existing paragraphs in place rather than
inserting new ones, to preserve the document's paragraph-index map.
"""
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers/p3-singapore/Manuscript_Blinded_MIR_2_revised.docx"


# =================================================================
# A. Section 4.5 (para 94) — extend with item-swap falsification
# =================================================================
P94_OLD = (
    "When DAI is reduced to a thinner, website-only measure, the "
    "moderation result weakens, suggesting that the richer baseline "
    "index draws meaningful explanatory variation from the electronic-"
    "payment indicators as well as from digital presence. By contrast, "
    "when the sample excludes micro-firms or is restricted to SMEs, "
    "the positive quadratic moderation pattern remains visible and in "
    "some cases becomes stronger. The exporters-only subsample is less "
    "stable because the number of exporting firms is small and the "
    "right tail remains thin, so those estimates should be described "
    "as suggestive rather than definitive."
)

# Add: item-swap falsification + drop-website sensitivity narrative
P94_NEW = (
    "When DAI is reduced to a thinner, website-only measure, the "
    "moderation result weakens, suggesting that the richer baseline "
    "index draws meaningful explanatory variation from the electronic-"
    "payment indicators as well as from digital presence. By contrast, "
    "when the sample excludes micro-firms or is restricted to SMEs, "
    "the positive quadratic moderation pattern remains visible and in "
    "some cases becomes stronger. The exporters-only subsample is less "
    "stable because the number of exporting firms is small and the "
    "right tail remains thin, so those estimates should be described "
    "as suggestive rather than definitive. "
    "An item-swap falsification test further supports the construct "
    "boundary between TCI and DAI: when the website indicator (c22b) "
    "is reassigned from DAI to TCI, the joint significance of the DAI "
    "moderation block collapses (joint F drops from 4.56 [p = .011] "
    "in the canonical specification to 1.88 [p = .154]), whereas "
    "swapping technology indicators in the opposite direction "
    "preserves the moderation pattern. The conditional-scaling "
    "interpretation therefore depends on the foundational digital-"
    "infrastructure indicators (website plus electronic payments) "
    "being correctly grouped under DAI rather than absorbed into a "
    "broader capability construct."
)


# =================================================================
# B. Section 7 limitation 3 (para 117) — expand IV suggestions
# =================================================================
P117_OLD = (
    "Panel data, successive enterprise-survey waves, linked "
    "administrative records, and stronger quasi-experimental designs "
    "— including instrumental-variable or difference-in-differences "
    "specifications keyed to digital-infrastructure rollouts or "
    "trade-policy shocks — would also help clarify whether the "
    "observed associations reflect scaling complementarities, "
    "selection effects, or other mechanisms."
)

P117_NEW = (
    "Panel data, successive enterprise-survey waves, linked "
    "administrative records, and stronger quasi-experimental designs "
    "— including instrumental-variable or difference-in-differences "
    "specifications keyed to digital-infrastructure rollouts or "
    "trade-policy shocks — would also help clarify whether the "
    "observed associations reflect scaling complementarities, "
    "selection effects, or other mechanisms. Concrete instrumental-"
    "variable strategies could exploit firm-to-port distance or local "
    "customs-clearance times as instruments for export intensity, and "
    "regional broadband-coverage shares or within-sector peer-adoption "
    "rates as instruments for foundational digital adoption, although "
    "Wolfolds and Siegel (2019) caution that such designs require both "
    "adequate first-stage strength and credible exclusion restrictions "
    "before they can deliver clean causal identification."
)


# =================================================================
# C. Section 5.1 implication 3 (para 105) — add TCE framing
# =================================================================
P105_OLD = (
    "Third, the DAI result suggests that foundational digital "
    "adoption is better understood as a conditional scaling resource "
    "than as a uniform firm-level productivity premium. Across much "
    "of the observed export-intensity distribution, the DAI "
    "association is weak or statistically indistinct, but it becomes "
    "more positive in the high-export tail, where firms face denser "
    "cross-border coordination and transaction demands. This pattern "
    "is consistent with the idea that Tier 1–2 digital adoption "
    "matters most when firms have sufficient international throughput "
    "to use digital interfaces and transaction-enabling systems "
    "intensively, while also underscoring that the evidence is "
    "concentrated in a thin-support region and should therefore be "
    "interpreted cautiously. The contribution here is thus not to "
    "establish a universal digitalization mechanism, but to clarify "
    "that foundational digital adoption and technological capability "
    "are analytically distinct and empirically associated with "
    "performance in different ways within this setting."
)

# Insert TCE sentence before the closing "The contribution here is thus..."
P105_NEW = (
    "Third, the DAI result suggests that foundational digital "
    "adoption is better understood as a conditional scaling resource "
    "than as a uniform firm-level productivity premium. Across much "
    "of the observed export-intensity distribution, the DAI "
    "association is weak or statistically indistinct, but it becomes "
    "more positive in the high-export tail, where firms face denser "
    "cross-border coordination and transaction demands. This pattern "
    "is consistent with the idea that Tier 1–2 digital adoption "
    "matters most when firms have sufficient international throughput "
    "to use digital interfaces and transaction-enabling systems "
    "intensively, while also underscoring that the evidence is "
    "concentrated in a thin-support region and should therefore be "
    "interpreted cautiously. From a transaction-cost perspective "
    "(Coase 1937; Williamson 1985), the conditional pattern is "
    "consistent with foundational digital interfaces absorbing "
    "coordination and information-processing costs that scale super-"
    "linearly with cross-border transaction volume; below the "
    "threshold at which firms operate at high export intensity, the "
    "marginal benefit of these systems may not exceed the fixed costs "
    "of installation and routine use, which would explain why no "
    "uniform productivity premium is observed across the full sample. "
    "The contribution here is thus not to establish a universal "
    "digitalization mechanism, but to clarify that foundational "
    "digital adoption and technological capability are analytically "
    "distinct and empirically associated with performance in "
    "different ways within this setting."
)


def replace_para(p, new_text):
    if not p.runs:
        p.text = new_text
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def main():
    doc = Document(DOC)

    print("[A] §4.5 (para 94): add item-swap falsification disclosure")
    cur = doc.paragraphs[94].text
    assert cur.strip() == P94_OLD.strip(), \
        f"Para 94 text mismatch — abort.\nGot: {cur[:200]}"
    replace_para(doc.paragraphs[94], P94_NEW)
    print(f"      Words: {len(P94_OLD.split())} → {len(P94_NEW.split())} "
          f"(+{len(P94_NEW.split()) - len(P94_OLD.split())})")

    print("\n[C] §5.1 implication 3 (para 105): add TCE framing")
    cur = doc.paragraphs[105].text
    assert cur.strip() == P105_OLD.strip(), \
        f"Para 105 text mismatch — abort.\nGot: {cur[:200]}"
    replace_para(doc.paragraphs[105], P105_NEW)
    print(f"      Words: {len(P105_OLD.split())} → {len(P105_NEW.split())} "
          f"(+{len(P105_NEW.split()) - len(P105_OLD.split())})")

    print("\n[B] §7 limitation 3 (para 117): expand IV suggestions")
    cur117 = doc.paragraphs[117].text
    assert P117_OLD in cur117, \
        f"Para 117 substring not found — abort.\nGot: {cur117[-300:]}"
    replace_para(doc.paragraphs[117], cur117.replace(P117_OLD, P117_NEW))
    new_count = len(doc.paragraphs[117].text.split())
    print(f"      Words: 137 → {new_count} (+{new_count - 137})")

    doc.save(DOC)
    print(f"\nSaved: {DOC}")


if __name__ == "__main__":
    main()
