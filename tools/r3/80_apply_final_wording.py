"""
R3 final wording pass — apply all editorial replacements per user's
ready-to-paste blocks across Abstract, Sections 1.2, 1.3, end of 2.1,
H2/H3/H4, Section 2.4, Section 3.3, Section 4.2, Section 4.4 (insert
clarification + soften closing).

Strategy:
  (1) In-place replacements (no paragraph count change) first.
  (2) Insertions (which add new paragraphs) last, walked from highest
      index to lowest so earlier indices remain stable.

This script is idempotent within a single run only — re-running on an
already-modified docx will replace text again but should not insert
duplicate paragraphs as long as we check by content sentinel before
inserting.
"""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers/p3-singapore/Manuscript_Blinded_MIR_2_revised.docx"


# =================================================================
# Wording blocks (verbatim from user's ready-to-paste text)
# =================================================================

ABSTRACT_P1 = (
    "This study examines how technological capability and digital "
    "adoption are associated with the internationalization–performance "
    "relationship among firms in Singapore, treated as an analytically "
    "informative, within-context setting of a digitally advanced economy. "
    "Using World Bank Enterprise Survey microdata for Singapore 2023, the "
    "analysis distinguishes a Technological Capability Index (TCI), "
    "capturing firm-internal capability depth, from a Digital Adoption "
    "Index (DAI), capturing foundational digital interfaces and "
    "transaction-enabling mechanisms."
)

ABSTRACT_P2 = (
    "Three findings emerge. First, within the observed export-intensity "
    "range, the internationalization–performance relationship is better "
    "characterized as predominantly positive with mild quadratic "
    "curvature than as a formally identified inverted-U. The fitted "
    "quadratic implies a turning point in the upper tail, but that point "
    "is imprecisely located and falls in a sparsely populated region of "
    "the data. Second, TCI is positively associated with labour "
    "productivity, while no statistically distinguishable moderation by "
    "TCI is detected under the present design. Third, DAI does not "
    "exhibit a large uniform productivity premium across firms; instead, "
    "its association with productivity becomes more positive at higher "
    "levels of export intensity, with the clearest signal concentrated "
    "in the high-export tail. Taken together, the findings provide "
    "within-context evidence from Singapore that sharpens the distinction "
    "between technological capability and foundational digital adoption "
    "and suggests that digital adoption may function more as a contingent "
    "scaling resource than as a uniform productivity advantage."
)

S12_P1 = (
    "Three gaps motivate this study. First, prior research has often "
    "bundled digitalization-related firm attributes into broad umbrella "
    "constructs, making it difficult to distinguish firm-internal "
    "technological capability from more basic forms of digital adoption. "
    "This matters because the two domains imply different mechanisms of "
    "advantage: technological capability concerns learning, innovation, "
    "and absorptive depth within the firm, whereas foundational digital "
    "adoption concerns the use of digital interfaces and transaction-"
    "enabling systems that may matter differently across stages of "
    "internationalization."
)

S12_P2 = (
    "Second, although work on digital internationalization has generated "
    "important conceptual advances, firm-level evidence remains limited "
    "on how basic digital adoption relates to export intensity within a "
    "single digitally advanced institutional environment. In particular, "
    "the literature still offers limited evidence on whether foundational "
    "digital tools matter uniformly across firms or become more relevant "
    "only when cross-border coordination demands intensify."
)

S12_P3 = (
    "Third, the nonlinear internationalization–performance literature "
    "has documented inverted-U patterns extensively, but the "
    "interpretation of such patterns in a digitally advanced setting "
    "remains under-specified. In a case such as Singapore, the "
    "unresolved issue is not whether one setting can establish a general "
    "boundary condition for all digitally mature economies, but how "
    "firm-level evidence from that setting should be interpreted "
    "relative to the conventional nonlinear literature. This study "
    "addresses that issue by examining whether the right-side decline "
    "of the curve is clearly identifiable within the export-intensity "
    "range actually occupied by firms in Singapore and by separating "
    "technological capability from foundational digital adoption in the "
    "same empirical framework."
)

S13_P1 = (
    "This study makes one primary contribution and two supporting "
    "contributions. Its primary contribution is to provide firm-level, "
    "within-context evidence from Singapore showing that the classic "
    "right-side decline of the internationalization–performance curve "
    "is not clearly identified within the export-intensity range "
    "occupied by most firms in the sample. In this sense, the study "
    "qualifies rather than overturns the conventional nonlinear "
    "literature by showing that, in this setting, the fitted quadratic "
    "is more defensibly read as predominantly positive with mild "
    "curvature than as a formally established inverted-U."
)

S13_P2 = (
    "The first supporting contribution is construct clarification. The "
    "analysis distinguishes technological capability from foundational "
    "digital adoption, thereby separating firm-internal capability depth "
    "rooted in learning, innovation, and technology absorption from more "
    "basic digitally enabled interfaces and transaction mechanisms. This "
    "distinction matters because the two constructs are associated with "
    "firm performance in different ways and should not be treated as "
    "interchangeable indicators of a single digital-capability domain."
)

S13_P3 = (
    "The second supporting contribution is substantive. The findings "
    "indicate that foundational digital adoption is better understood as "
    "a conditional scaling resource than as a uniform productivity "
    "premium: its association with productivity is weak across much of "
    "the export-intensity distribution and becomes more positive only "
    "in the high-export tail. This does not establish a universal "
    "digital-frontier mechanism, but it does provide within-context "
    "evidence that the productivity relevance of foundational digital "
    "adoption depends on the intensity of cross-border operations."
)

S21_END = (
    "An important implication of the nonlinear I–P literature is that "
    "the visibility of an inverted-U depends not only on theory but "
    "also on where firms actually lie in the observed distribution of "
    "internationalization intensity. If firms operate in a setting "
    "where digital infrastructure and transaction systems reduce some "
    "coordination frictions, the right-side decline may become harder "
    "to detect within the export range most firms occupy, even if "
    "coordination costs do not disappear in principle. Singapore is "
    "therefore analytically useful not because it can by itself "
    "establish a general boundary condition for digitally mature "
    "economies, but because it provides a within-context case for "
    "examining how the conventional I–P logic appears when firms "
    "operate in a digitally advanced institutional environment."
)

H2_REASONING = (
    "A subtler prediction concerns whether technological capability "
    "alters the shape of the internationalization–performance "
    "relationship rather than simply raising average productivity. From "
    "a non-location-bound firm-specific advantage perspective, "
    "technological capability may travel across markets without "
    "requiring the same degree of adaptation as context-dependent "
    "coordination tools. This suggests that its most stable empirical "
    "role may lie in its direct association with productivity, while "
    "any moderation of the FSTS–performance relationship is less "
    "certain and should be treated as an empirical question rather than "
    "presumed ex ante."
)
H2_STATEMENT = (
    "Hypothesis 2 (H2). Technological capability (TCI) is positively "
    "associated with labour productivity in Singapore, while its "
    "moderation of the internationalization–performance relationship "
    "is treated as an open empirical question."
)

H3_REASONING = (
    "Digital adoption may improve firm performance through faster "
    "information flows, lower transaction frictions, and more efficient "
    "digital exchange, but such benefits are unlikely to be uniform "
    "across firms. In the present empirical setting, the available "
    "indicators capture foundational digital interfaces and transaction-"
    "enabling mechanisms rather than deeper digitally integrated "
    "organizational capabilities. Their productivity relevance is "
    "therefore not expected to appear as a large, unconditional firm-"
    "wide premium, but as an association whose strength may depend on "
    "the scale at which firms actually use those digital channels."
)
H3_STATEMENT = (
    "Hypothesis 3 (H3). The productivity association of digital "
    "adoption (DAI) in Singapore is conditional rather than uniform "
    "across firms."
)

H4_REASONING = (
    "The most specific prediction concerns whether foundational digital "
    "adoption becomes more relevant when firms face denser cross-border "
    "coordination demands. Firms with higher export intensity process "
    "more transactions across customers, suppliers, and markets, making "
    "digital interfaces and electronic transaction systems potentially "
    "more valuable as scaling mechanisms. By contrast, firms with low "
    "export intensity may have fewer opportunities to translate basic "
    "digital adoption into measurable labour-productivity gains."
)
H4_STATEMENT = (
    "Hypothesis 4 (H4). The association between digital adoption (DAI) "
    "and firm performance becomes more positive at higher levels of "
    "export intensity in Singapore."
)

S24_P1 = (
    "Figure 1 summarizes the conceptual model. The model treats TCI and "
    "DAI as analytically distinct firm-level constructs that are "
    "associated with firm performance through different channels. TCI "
    "is expected to show a more stable direct association with labour "
    "productivity than a clearly identified moderation pattern, "
    "reflecting its role as firm-internal capability depth grounded in "
    "learning, innovation, and technology absorption. DAI, by contrast, "
    "is theorized to matter contingently through export intensity rather "
    "than as a uniform direct premium, because the observed indicators "
    "capture foundational digital interfaces and transaction-enabling "
    "mechanisms whose relevance depends on how intensively firms engage "
    "in cross-border operations. In this sense, H3 and H4 are jointly "
    "informative: H3 concerns the non-uniformity of the DAI–productivity "
    "association, whereas H4 specifies the direction of that "
    "contingency across export intensity. The model is therefore "
    "positioned as a within-context framework for interpreting firm-"
    "level evidence from Singapore rather than as a design that can, on "
    "its own, establish a general boundary condition for digitally "
    "mature economies."
)

S33_P1 = (
    "The empirical analysis uses ordinary least squares with broad "
    "sector fixed effects and HC1 heteroskedasticity-robust standard "
    "errors throughout. Models are estimated sequentially, beginning "
    "with a controls-only specification and then moving to a linear "
    "internationalization model, a quadratic full-sample benchmark, "
    "direct-effect models for TCI and DAI, and full specifications "
    "that incorporate moderation terms. In light of the strong "
    "concentration of firms at zero exports, the baseline quadratic "
    "specification is interpreted primarily as a descriptive full-"
    "sample benchmark rather than as a standalone structural test of "
    "the conventional inverted-U logic. This is important because the "
    "full-sample polynomial is necessarily influenced by the domestic-"
    "versus-exporter split, whereas the theoretical question developed "
    "in Sections 1–2 concerns how performance evolves across the "
    "continuous export-intensity range. The later specifications "
    "therefore focus on whether TCI and DAI exhibit distinct direct "
    "and contingent associations, and the interpretation of nonlinear "
    "structure is supplemented by extensive–intensive reasoning and "
    "support-aware diagnostics rather than inferred from the quadratic "
    "fit alone. Accordingly, the analysis does not treat the full-"
    "sample polynomial as sufficient evidence of a formally identified "
    "inverted-U in the observed data range."
)

S42_P1 = (
    "In the baseline quadratic specification (Model M2), the linear "
    "FSTS term is positive and statistically significant, whereas the "
    "squared term is negative and only marginally significant. The "
    "fitted curve therefore implies a turning point in the upper tail "
    "of the export-intensity distribution, near FSTS = 82% on the "
    "original scale, but that point is imprecisely located and falls in "
    "a sparsely populated region of the data. The 5,000-replication "
    "bootstrap recovers an inverted-U shape frequently (96.3% of "
    "replications), yet the corresponding 95% percentile confidence "
    "interval for the turning point is wide ([53%, 253%]), and the "
    "Lind–Mehlum test does not formally confirm a conventional "
    "inverted-U at conventional significance levels (p = .303). The "
    "most defensible interpretation is therefore not that the inverted-"
    "U is firmly established, but that the full-sample relationship is "
    "predominantly positive with mild quadratic curvature over the "
    "observed export-intensity range. In this sense, the quadratic fit "
    "is informative descriptively, while stronger claims about a "
    "structurally identified right-side decline remain beyond what the "
    "present data can support. Read this way, the evidence qualifies "
    "rather than overturns the conventional nonlinear literature, "
    "while keeping interpretation within the empirical support actually "
    "available in the Singapore sample."
)

# Section 4.2 had 2 content paragraphs (69 + 70). New version compresses
# both into a single paragraph (S42_P1). Para 70 will be cleared.
S42_P2_CLEAR = ""

# Section 4.4 closing softening (replaces para 86)
S44_CLOSE = (
    "In this setting, foundational digital adoption is better "
    "interpreted as a conditional scaling resource whose productivity "
    "relevance becomes more visible at higher export intensity, while "
    "remaining weak or statistically indistinct across much of the "
    "distribution."
)

# Section 4.4 clarification — to be INSERTED after para 81
S44_CLARIFY = (
    "One clarification is important for interpretation. Because FSTS is "
    "mean-centered before squaring and the interaction terms are "
    "constructed from the centered measure, the coefficient on DAI in "
    "Model M8 is not numerically equivalent to the marginal effect of "
    "DAI at FSTS = 0 on the original scale. For that reason, the near-"
    "zero direct DAI coefficient in M8 should not be read as "
    "contradicting the small positive marginal effect reported for "
    "domestic firms in Table 4. The more relevant substantive pattern "
    "is that the DAI association is weak across much of the "
    "distribution and becomes more positive only in the high-export "
    "tail, where cross-border coordination demands are denser but "
    "empirical support is also thinner."
)

# Caption tweak for Figure 1 (para 44) — minor: replace closing phrase
F1_CAPTION_OLD_PHRASE = ("The setting is Singapore as an extreme-case, "
                         "within-context observation of a digitally "
                         "mature economy")
F1_CAPTION_NEW_PHRASE = ("The setting is Singapore as a within-context, "
                         "analytically informative case of a digitally "
                         "advanced economy")


# =================================================================
# Helpers
# =================================================================
def replace_para(p, new_text: str):
    """Replace paragraph text, preserving the first run's formatting."""
    if not p.runs:
        p.text = new_text
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def insert_after(template_p, new_text: str) -> None:
    """Insert a new paragraph (cloned from template) right AFTER template_p."""
    new_xml = deepcopy(template_p._p)
    template_p._p.addnext(new_xml)
    new_p = Paragraph(new_xml, template_p._parent)
    replace_para(new_p, new_text)


def main():
    doc = Document(DOC)

    # ============================================================
    # 1. In-place replacements
    # ============================================================
    print("[1] Replacing paragraphs in place...")
    replacements = [
        (2,  ABSTRACT_P1,      "Abstract para 1"),
        (3,  ABSTRACT_P2,      "Abstract para 2"),
        (12, S12_P1,           "Section 1.2 para 1"),
        (13, S12_P2,           "Section 1.2 para 2"),
        (14, S12_P3,           "Section 1.2 para 3"),
        (16, S13_P1,           "Section 1.3 para 1"),
        (17, S13_P2,           "Section 1.3 para 2"),
        (23, S21_END,          "Section 2.1 closing"),
        (33, H2_REASONING,     "H2 reasoning"),
        (34, H2_STATEMENT,     "H2 statement"),
        (36, H3_REASONING,     "H3 reasoning"),
        (37, H3_STATEMENT,     "H3 statement"),
        (39, H4_REASONING,     "H4 reasoning"),
        (40, H4_STATEMENT,     "H4 statement"),
        (42, S24_P1,           "Section 2.4 para 1"),
        (58, S33_P1,           "Section 3.3 para 1"),
        (61, "",               "Section 3.3 para 4 (subsumed by new para 1)"),
        (69, S42_P1,           "Section 4.2 para 1"),
        (70, S42_P2_CLEAR,     "Section 4.2 para 2 (cleared)"),
        (86, S44_CLOSE,        "Section 4.4 closing softened"),
    ]
    for idx, text, label in replacements:
        replace_para(doc.paragraphs[idx], text)
        print(f"      ✓ [{idx:3}] {label}")

    # ============================================================
    # 2. Caption tweak (para 44 — Figure 1 caption ending)
    # ============================================================
    print("\n[2] Tweaking Figure 1 caption (para 44)...")
    cap = doc.paragraphs[44]
    if F1_CAPTION_OLD_PHRASE in cap.text:
        new_text = cap.text.replace(F1_CAPTION_OLD_PHRASE,
                                    F1_CAPTION_NEW_PHRASE)
        replace_para(cap, new_text)
        print(f"      ✓ Caption phrase updated")
    else:
        print(f"      (no match for old phrase — caption may have changed)")

    # ============================================================
    # 3. Insertions (highest index first to keep earlier indices stable)
    # ============================================================
    print("\n[3] Inserting new paragraphs (reverse-order to preserve indices)...")

    # 3a. Insert Section 4.4 clarification AFTER para 81 (M8 paragraph)
    sentinel_clarify = "One clarification is important for interpretation"
    if sentinel_clarify not in "\n".join(p.text for p in doc.paragraphs):
        insert_after(doc.paragraphs[81], S44_CLARIFY)
        print(f"      ✓ Inserted Section 4.4 clarification after para 81")
    else:
        print(f"      (clarification already present — skip)")

    # 3b. Insert Section 1.3 third paragraph AFTER para 17
    # (this is the "second supporting contribution" paragraph)
    sentinel_s13_p3 = "The second supporting contribution is substantive"
    if sentinel_s13_p3 not in "\n".join(p.text for p in doc.paragraphs):
        insert_after(doc.paragraphs[17], S13_P3)
        print(f"      ✓ Inserted Section 1.3 third paragraph after para 17")
    else:
        # Need to also UPDATE existing — find it and replace
        for i, p in enumerate(doc.paragraphs):
            if sentinel_s13_p3 in p.text:
                replace_para(p, S13_P3)
                print(f"      ✓ Updated existing Section 1.3 third "
                      f"paragraph at index {i}")
                break

    doc.save(DOC)
    print(f"\nSaved: {DOC}")


if __name__ == "__main__":
    main()
