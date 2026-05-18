"""
R3 Phase 0 — Apply corrected Table 3 to the manuscript.

Replaces the contents of Table 3 in
`papers/p3-singapore/Manuscript_Blinded_MIR_2_revised.docx` with the
canonical re-estimation produced by `02_audit_table3.py`.

The replacement preserves the table structure (7 rows × 6 columns) and
only updates cell text. Header row (row 0) is unchanged.
"""
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers" / "p3-singapore" / "Manuscript_Blinded_MIR_2_revised.docx"

CORRECTED_ROWS = [
    # (Specification, N, TCI β_z, FSTS² × DAI, Joint F (p), Adj. R²)
    ("Baseline (TCI_full + DAI_rich)",   "617", "0.153***", "3.119**",  "4.56 (.011)", "0.196"),
    ("R1: DAI_thin (website c22b only)", "623", "0.180***", "1.552",    "4.01 (.019)", "0.188"),
    ("R2: TCI_thin (e6 + b8)",           "617", "0.159***", "2.930**",  "4.27 (.014)", "0.199"),
    ("R3: Excl micro-firms (<10 empl)",  "464", "0.170***", "3.521**",  "4.61 (.010)", "0.219"),
    ("R4: SMEs only (≤200 empl)",        "595", "0.156***", "3.505**",  "5.30 (.005)", "0.199"),
    ("R5: Exporters only (FSTS > 0)",     "84", "0.130",    "2.821**",  "6.32 (.003)", "0.165"),
]


def _set_cell_text(cell, text: str):
    """Replace cell text while preserving the first run's formatting."""
    if cell.paragraphs and cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = text
        # Clear any extra runs
        for r in cell.paragraphs[0].runs[1:]:
            r.text = ""
    else:
        cell.text = text


def main():
    doc = Document(DOC)
    # Table 3 is index 3 in the document's table list (verified manually).
    table = doc.tables[3]
    assert len(table.rows) == 7, f"Expected 7 rows, got {len(table.rows)}"
    assert len(table.rows[0].cells) == 6, "Expected 6 columns"

    print("Before fix:")
    for ri, row in enumerate(table.rows):
        print(f"  R{ri}: {[c.text.strip()[:40] for c in row.cells]}")

    for i, vals in enumerate(CORRECTED_ROWS, start=1):
        for j, v in enumerate(vals):
            _set_cell_text(table.rows[i].cells[j], v)

    print("\nAfter fix:")
    for ri, row in enumerate(table.rows):
        print(f"  R{ri}: {[c.text.strip()[:40] for c in row.cells]}")

    doc.save(DOC)
    print(f"\nSaved: {DOC}")


if __name__ == "__main__":
    main()
