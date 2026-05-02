"""
R3 Phase 4 — Apply manuscript rewrites for framing, hypotheses, Section 7.

Rewrites in this script:
  - Title (paragraph 0): drop "Digital-Frontier Economy", recast.
  - Abstract second paragraph (3): align with Phase 1–3 evidence.
  - H2 (para 34): non-sig interaction != absence of moderation.
  - H3 (para 37): drop "positive on average"; conditional logic.
  - Section 7 first three caveats (113, 114, 115): D2 fix; replace
    adj-R² rationale with thin-tail + small-subsample logic; drop the
    digital-frontier boundary-condition claim.

Sync also:
  - Cover_Letter.docx: title and framing.
  - Title_Page.docx: title.
"""
from pathlib import Path
from docx import Document

REPO = Path(__file__).resolve().parents[2]
DOC = REPO / "papers" / "p3-singapore" / "Manuscript_Blinded_MIR_2_revised.docx"
COVER = REPO / "papers" / "p3-singapore" / "Cover_Letter.docx"
TITLE_PAGE = REPO / "papers" / "p3-singapore" / "Title_Page.docx"

NEW_TITLE = ("Technological Capability, Digital Adoption, and the "
             "Internationalization–Performance Relationship: A Firm-Level "
             "Study of Singapore")

NEW_ABSTRACT_PARA2 = (
    "Three findings emerge. First, the full-sample internationalization–"
    "performance pattern shows an inverted-U shape that is robust in shape "
    "(96.3% of bootstrap replications recover an inverted-U) but only "
    "loosely identified in location: the implied turning point sits at "
    "82% export intensity, with a 95% percentile bootstrap confidence "
    "interval of [53%, 253%]. Within an extensive–intensive split design "
    "that addresses the dominance of zero-FSTS firms (82% of the sample), "
    "the curvature is not detectable inside the exporter subsample alone, "
    "indicating that the full-sample polynomial is best read as a "
    "descriptive baseline rather than as a structurally identified curve. "
    "Second, technological capability (TCI) is positively and "
    "consistently associated with labour productivity across all "
    "specifications and across leave-one-out and trimmed-tail "
    "perturbations; under the present design and statistical power, no "
    "distinguishable moderation of the FSTS–productivity relationship by "
    "TCI is detected. Third, digital adoption (DAI) shows a conditional "
    "association with productivity rather than a uniform firm-level "
    "premium: its productivity association becomes more positive at "
    "higher levels of export intensity, and this moderation pattern "
    "survives leave-one-out re-estimation, item-swap falsification, and "
    "drop-FSTS-above-70% trimming. The study contributes refined "
    "boundary conditions for the internationalization–performance "
    "literature in Singapore as an extreme-case, within-context setting, "
    "and clarifies that, in a measurement architecture that distinguishes "
    "technological capability from foundational digital adoption, the "
    "two constructs play analytically different roles: TCI shifts the "
    "productivity level, while DAI complements export-intensity scaling.")

NEW_H2 = (
    "Hypothesis 2 (H2). Technological capability (TCI) is positively "
    "associated with labour productivity in Singapore. Whether TCI also "
    "moderates the curvature of the internationalization–performance "
    "relationship is treated as an open empirical question; the present "
    "design tests for, but does not presume, a moderation channel.")

NEW_H3 = (
    "Hypothesis 3 (H3). The productivity association of digital adoption "
    "(DAI) in Singapore is conditional on export intensity rather than "
    "uniform across firms. Specifically, foundational digital adoption "
    "is expected to complement the productivity returns to higher levels "
    "of export intensity; the moderation pattern is formally tested in "
    "H4.")

NEW_SEC7_PARA1 = (
    "Three limitations warrant explicit caution in interpreting the "
    "findings. First, the analysis is a single-country cross-section, "
    "and the identification strategy is therefore associational rather "
    "than causal. The observed positive associations among TCI, DAI, "
    "internationalization, and labour productivity admit at least two "
    "non-causal interpretations consistent with the data. The first is "
    "productivity-driven self-selection into digital adoption: more "
    "productive firms may have the financial slack and managerial "
    "bandwidth to install electronic-payment systems and digital "
    "interfaces, so the observed DAI–productivity association may "
    "reflect this selection channel rather than a productivity-enhancing "
    "effect of digital adoption per se. The second is productivity-"
    "driven self-selection into international markets, in which only "
    "firms above a productivity threshold profitably overcome the fixed "
    "costs of exporting; the observed FSTS–DAI interaction could "
    "therefore reflect productivity-induced co-movement of export "
    "intensity and digital-system utilization rather than a digital-"
    "system-induced productivity gain. Standard ex post diagnostics such "
    "as non-significant inverse Mills ratios from Heckman-style models "
    "do not constitute affirmative evidence of unbiasedness in the "
    "absence of a valid exclusion restriction (Wolfolds & Siegel 2019). "
    "The disciplined response is therefore to read the present results "
    "as extreme-case, within-context evidence from Singapore: the "
    "estimated patterns describe how productivity, technological "
    "capability, and digital adoption co-move with export intensity in "
    "a digitally mature economy, while leaving causal identification of "
    "any specific mechanism to designs that cannot be implemented in "
    "the present study.")

NEW_SEC7_PARA2 = (
    "Second, the sample contains a thin right tail of high-intensity "
    "exporters: 82.2% of firms report zero exports, 4.5% exceed 50% "
    "export intensity, and only 3.2% exceed 70%. Two consequences "
    "follow. First, the implied quadratic turning point at 82% on the "
    "FSTS scale is identified from a sparsely populated region; "
    "5,000-replication cluster bootstrap places the 95% percentile "
    "confidence interval at [53%, 253%], so the turning point is best "
    "read as an indicative descriptive feature rather than a "
    "structurally identified inflection — the inverted-U *shape* is "
    "robust (recovered in 96.3% of bootstrap replications) but its "
    "*location* is imprecise. Second, the DAI moderation pattern attains "
    "conventional joint significance in the full-sample, leave-one-out "
    "(100% of LOO fits), and trimmed-tail re-estimations (drop "
    "FSTS > 70%: F = 5.51, p = .004; drop FSTS > 80%: F = 4.42, "
    "p = .012); within the exporter-only subsample (N = 84) the joint "
    "F-test for DAI moderation remains significant (F = 6.32, p = .003) "
    "but the 95% confidence intervals for individual coefficients are "
    "wider, reflecting the small subsample size rather than a fragile "
    "underlying pattern. The appropriate caution therefore rests on "
    "(i) the thin upper-tail support for the precise location of the "
    "turning point and (ii) the borderline precision of individual "
    "moderation coefficients in the small exporter subsample, not on "
    "differences in adjusted R² across samples of unequal size and "
    "composition.")

NEW_SEC7_PARA3 = (
    "Third, the estimated moderation effect is statistically detectable "
    "but substantively modest, and the present results should be read "
    "as an initial firm-level test of a conditional digital-adoption "
    "argument in one digitally mature setting rather than as definitive "
    "evidence of a generalized boundary condition. The DAI indicators "
    "available in WBES 2023 capture Tier 1–2 digital adoption — digital "
    "presence and digital transaction usage — rather than Tier 3–4 "
    "digital capability or dynamic capability; indicator-sensitivity "
    "analysis shows that the moderation pattern depends on the "
    "foundational digital-infrastructure indicators (firm website and "
    "electronic-payment-to-suppliers usage) and weakens when these "
    "specific indicators are dropped. Although equivalence-testing "
    "approaches (Lakens et al. 2018) could in principle quantify the "
    "support for null moderation in subsamples, the present sample size "
    "limits their power, and the more cautious reading is that absence "
    "of moderation cannot be inferred from non-significance alone. "
    "Whether the scale-enabling-complement pattern persists, attenuates, "
    "or is amplified under richer measures of digitally integrated "
    "organizational capabilities remains an open empirical question. "
    "Future research should therefore examine whether the same pattern "
    "appears in other digitally mature economies, such as Hong Kong, "
    "the Republic of Korea, and Taiwan, and under purpose-designed "
    "instruments that capture deeper digital dynamic capabilities. "
    "Future research should also extend this measurement architecture "
    "by linking WBES-type microdata, where confidentiality arrangements "
    "permit, to firm-level databases or by using purpose-built survey "
    "instruments that better capture process-level digital integration "
    "and higher-order capability depth. More broadly, future work "
    "should examine whether the scale-enabling-complement pattern "
    "operates through identifiable mechanisms such as supplier-side "
    "digital integration, customer-facing digital channels, or internal "
    "ERP integration. Causal identification, in particular, would "
    "benefit from firm-level panel data, successive enterprise-survey "
    "waves, national statistical-office registers, or proprietary "
    "databases that permit firm fixed-effects estimation. Where "
    "suitable instruments are available, such as regional digital-"
    "infrastructure rollouts or plausibly exogenous trade-policy "
    "shocks, instrumental-variable or difference-in-differences designs "
    "would further upgrade the present associational evidence to "
    "causal identification.")


def replace_paragraph(p, new_text):
    """Replace text of a paragraph while preserving the first run's style."""
    if not p.runs:
        p.text = new_text
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def main():
    print("=" * 72)
    print("R3 Phase 4 — Apply manuscript rewrites")
    print("=" * 72)

    doc = Document(DOC)
    paragraphs = doc.paragraphs

    # 1. Title (para 0)
    print(f"\n[1] Title (para 0)")
    print(f"    OLD: {paragraphs[0].text[:100]}...")
    replace_paragraph(paragraphs[0], NEW_TITLE)
    print(f"    NEW: {NEW_TITLE[:100]}...")

    # 2. Abstract second paragraph (3)
    print(f"\n[2] Abstract findings (para 3)")
    replace_paragraph(paragraphs[3], NEW_ABSTRACT_PARA2)
    print(f"    Replaced ({len(NEW_ABSTRACT_PARA2)} chars).")

    # 3. H2 (para 34)
    print(f"\n[3] H2 (para 34)")
    replace_paragraph(paragraphs[34], NEW_H2)

    # 4. H3 (para 37)
    print(f"\n[4] H3 (para 37)")
    replace_paragraph(paragraphs[37], NEW_H3)

    # 5. Section 7 paragraphs (113, 114, 115)
    print(f"\n[5] Section 7 paragraph 1 (para 113) — extreme-case framing")
    replace_paragraph(paragraphs[113], NEW_SEC7_PARA1)

    print(f"[6] Section 7 paragraph 2 (para 114) — D2 fix")
    replace_paragraph(paragraphs[114], NEW_SEC7_PARA2)

    print(f"[7] Section 7 paragraph 3 (para 115) — soft-pedal claims")
    replace_paragraph(paragraphs[115], NEW_SEC7_PARA3)

    # 6. Also fix the abstract first paragraph: drop "digital-frontier economy"
    para2 = paragraphs[2].text
    para2_new = para2.replace(
        "in a digital-frontier economy",
        "in Singapore, an extreme-case, within-context setting of a "
        "digitally mature economy"
    ).replace(
        "in a digital-frontier",
        "in an extreme-case Singapore"
    )
    if para2_new != para2:
        replace_paragraph(paragraphs[2], para2_new)
        print(f"[8] Abstract para 1 (para 2) — recast 'digital-frontier'")

    doc.save(DOC)
    print(f"\nSaved: {DOC}")

    # ---- Sync Cover Letter ----
    print("\n" + "=" * 72)
    print("Sync Cover Letter")
    print("=" * 72)
    cover = Document(COVER)
    n_changed = 0
    for p in cover.paragraphs:
        old = p.text
        new = old
        # Update title quote inside the cover letter
        new = new.replace(
            "Technological Capability, Digital Adoption, and the "
            "Internationalization–Performance Relationship in a "
            "Digital-Frontier Economy: Evidence From Singapore",
            NEW_TITLE
        )
        # Strip remaining "digital-frontier" claims
        new = new.replace(
            "digital-frontier environment", "extreme-case Singapore setting"
        ).replace(
            "digital-frontier economy", "extreme-case Singapore setting"
        ).replace(
            "in a digital-frontier", "in an extreme-case Singapore"
        )
        if new != old:
            replace_paragraph(p, new)
            n_changed += 1
    cover.save(COVER)
    print(f"  Updated {n_changed} paragraphs in Cover Letter.")

    # ---- Sync Title Page ----
    print("\nSync Title Page")
    title_doc = Document(TITLE_PAGE)
    n_changed = 0
    for p in title_doc.paragraphs:
        old = p.text
        if "Digital-Frontier Economy" in old or "Digital-Frontier" in old:
            new = NEW_TITLE if old.startswith("Technological Capability") else \
                  old.replace("Digital-Frontier Economy",
                              "Singapore Firm-Level Study")
            replace_paragraph(p, new)
            n_changed += 1
    title_doc.save(TITLE_PAGE)
    print(f"  Updated {n_changed} paragraphs in Title Page.")

    print("\nPhase 4 rewrites applied.")


if __name__ == "__main__":
    main()
