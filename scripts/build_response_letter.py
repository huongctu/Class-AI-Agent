"""Build the point-by-point response letter for the IJoEM reviewer report.

The letter maps every overall and detailed reviewer comment from the v5.6
review round to the specific revisions implemented in v5.7. It can be
attached alongside the cover letter and manuscript at submission.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm

OUT = Path("/home/user/Class-AI-Agent/response_letter_to_reviewer.docx")

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)


def H1(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(13)


def H2(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(11)


def P(t, italic=False, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.italic = italic
    r.bold = bold


def Q(t):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    r = p.add_run(t)
    r.italic = True


# ---------------------------------------------------------------------------
H1("Response to Reviewer Comments")
P("Manuscript ID: [to be assigned]")
P('Title: "Revisiting the Internationalisation–Performance Relationship in '
  'an Emerging Market: The Roles of Technological Capability and Digital '
  'Adoption"')
P("Journal: International Journal of Emerging Markets (Emerald)")
P("Submission round: Revised manuscript v5.7 (responding to v5.6 review)")
P("")

P("Dear Editor and Reviewer,")
P("")
P("We thank the Reviewer for the constructive and substantive comments on "
  "our manuscript. The comments have materially improved both the framing "
  "and the empirical evidence of the paper. Below we provide a point-by-"
  "point response, mapping each comment to the specific revisions in v5.7. "
  "All revisions are within the scope of the harmonised World Bank Enterprise "
  "Survey (WBES) data for Vietnam (waves 2009 / 2015 / 2023). Section "
  "references below are to v5.7.")
P("")

# ---------------------------------------------------------------------------
H1("Part A — Overall Feedback")

# A1 ------------------------------------------------------------------------
H2("A1. Digital adoption construct (over-claim relative to the WBES measure)")
P("Reviewer comment.", italic=True)
Q('"The primary DAI_z is a within-wave standardised website-presence dummy, '
  'while the abstract, §§2.3–2.4, §§5.1–5.5, and H3/H4 often speak in '
  'broader terms: digital adoption, transaction support, digital '
  'capability, and sometimes a digital-transformation channel."')
P("Our response.", bold=True)
P("We agree. We have re-theorised the primary DAI_z explicitly and "
  "uniformly across the manuscript as website-based digital presence / "
  "foundational website adoption. Specifically:")
P("• §2.3 has been retitled to \"Website-based digital presence and firm "
  "performance\" and now opens with: \"The primary DAI_z used in this "
  "paper is a website-based digital presence measure: a binary indicator "
  "of whether the firm has its own website. This is a foundational and "
  "cross-wave-comparable marker of digital adoption — it does not measure "
  "transaction-level digital integration, electronic payment "
  "infrastructure, or digital transformation in the Bharadwaj et al. "
  "(2013) / Verhoef et al. (2021) / Vial (2019) sense.\"")
P("• §3.2 now reads, unambiguously: \"The primary Digital Adoption Index "
  "(DAI_z) is the within-wave standardised website-presence indicator c22b… "
  "Under this revised primary specification, no item is shared between the "
  "TCI and DAI composites — e6 belongs exclusively to the capability "
  "construct, while c22b alone serves as a harmonised cross-wave proxy for "
  "basic digital presence.\"")
P("• The abstract and the structured Highlights have been rewritten to "
  "use \"website-based digital-presence indicator\" / \"foundational "
  "website adoption\" rather than \"digital transformation\" or "
  "\"transaction support\".")
P("• The 2023 e-payment items remain only inside §4.5 Panel B (DAI_rich) "
  "as a clearly exploratory 2023-only extension, not as validation of the "
  "harmonised cross-wave construct.")
P("• A new §5.2.1 (\"An alternative reading: proxy obsolescence\") names "
  "explicitly the possibility that the falling DAI direct association "
  "across waves partly reflects the shrinking informational content of "
  "the c22b indicator (a website was distinguishing in 2009 and is "
  "routine by 2023) rather than only a structural change in the payoff. "
  "We do not treat proxy obsolescence and stage contingency as mutually "
  "exclusive, and flag this because it is not testable within the cross-"
  "wave-comparable WBES instrument.")
P("")

# A2 ------------------------------------------------------------------------
H2("A2. Technological capability construct (label too broad relative to b8 + e6)")
P("Reviewer comment.", italic=True)
Q('"The paper\'s technological-capability argument is theoretically rich, '
  'but the primary measure is relatively thin… Section 3.2 operationalises '
  'the primary TCI_z as the mean of internationally recognised quality '
  'certification and foreign-licensed technology. Those two items… may '
  'capture standards compliance, foreign-market linkage, or exposure to '
  'foreign technology as much as firm-internal capability stocks."')
P("Our response.", bold=True)
P("We agree, and have narrowed the label and theoretical framing of the "
  "primary construct accordingly:")
P("• §2.2 has been retitled to \"Foreign-technology and standards "
  "capability and firm performance\" and now states explicitly: \"this "
  "paper uses a measurement-tight reading of technological capability: a "
  "foreign-technology and standards capability that captures a firm's "
  "exposure to externally validated technological inputs… This is one "
  "observable facet of the broader Cohen-Levinthal (1990) absorptive-"
  "capacity construct and the dynamic-capability construct of Teece "
  "(2007), but the present measure does not claim to identify the full "
  "absorptive-capacity stock; it captures the externally facing component "
  "of that stock.\"")
P("• H2 has been rewritten to: \"Foreign-technology / standards "
  "capability (TCI_z) is positively associated with firm performance in "
  "Vietnam.\"")
P("• §4.5 Panel A (TCI_full, augmented with product innovation h1 and "
  "R&D h8) is now framed explicitly as a boundary condition rather than "
  "as validation of the thinner measure: \"if direct effects attenuate "
  "when innovation items are added, this indicates that the primary "
  "measure is informative specifically about the foreign-technology / "
  "standards channel, not the broader absorptive-capacity stock.\" The "
  "2015 attenuation under TCI_full is reported and interpreted as such.")
P("• The Originality / Value statement and the Discussion (§5.2, §5.5) "
  "have been edited to refer to \"foreign-technology / standards "
  "capability\" rather than to \"deeper capability stocks\" or "
  "\"absorptive capacity\" when describing the primary construct.")
P("")

# A3 ------------------------------------------------------------------------
H2("A3. Cross-wave / lifecycle interpretation runs ahead of the formal evidence")
P("Reviewer comment.", italic=True)
Q('"Panel F shows that all FSTS_c, FSTS_c², FSTS_c × DAI_z, and FSTS_c² × '
  'DAI_z differences are non-significant. The clearest formal cross-wave '
  'evidence concerns the direct DAI_z coefficient… If the paper wants to '
  'retain a strong lifecycle claim, it would help to add common-'
  'specification evidence that directly tests cross-wave differences: '
  'pooled wave-by-focal-variable interactions…"')
P("Our response.", bold=True)
P("We have implemented both the formal test and the recalibration:")
P("• A new §4.5 Panel I (\"Pooled wave × focal interaction test\") "
  "estimates the pooled M8 with a saturated set of wave interactions on "
  "every focal term (FSTS_c × wave, FSTS_c² × wave, DAI_z × wave, TCI_z × "
  "wave). The joint Wald tests confirm the descriptive Panel F result: "
  "only DAI_z × wave is statistically distinguishable from the pooled "
  "average (joint p = .016); the FSTS_c, FSTS_c², TCI_z direct-effect "
  "cross-wave differences are not separable (all joint p > .25); and the "
  "FSTS_c × DAI_z and FSTS_c² × DAI_z cross-wave differences are not "
  "separable either (joint p > .55).")
P("• Section 4.1 and Section 5 lifecycle wording is recalibrated "
  "throughout from \"stage contingency\" to \"wave-specific associations "
  "consistent with stage contingency\". The §4.1 closing now reads: "
  "\"only the DAI direct shifts are cross-wave-distinguishable; the FSTS "
  "curvature and the FSTS × DAI moderation differences across waves are "
  "not statistically separable in the pooled saturated specification. We "
  "therefore read the pattern as wave-specific associations consistent "
  "with stage contingency, not as a fully cross-wave-identified "
  "structural shift.\"")
P("• The abstract, Highlights and §5.5 (Policy implications) have been "
  "edited to reflect the same calibration.")
P("")

# A4 ------------------------------------------------------------------------
H2("A4. Inverted-U could be driven by the zero-inflated FSTS distribution")
P("Reviewer comment.", italic=True)
Q('"Exporter shares are 28.4 %, 20.8 % and 18.8 % across the three waves. '
  'That means the quadratic in FSTS is being estimated on a bounded, zero-'
  'inflated variable with a large participation margin at zero… Exporter-'
  'only estimates… and distributional evidence around the 39–46 % turning '
  'points would all help."')
P("Our response.", bold=True)
P("We have added two new diagnostics and adjusted the H1 wording:")
P("• New §4.5 Panel H (\"Exporter-only sub-sample, FSTS > 0\"). We re-fit "
  "M2 / M7 / M8 on the exporters only (N ≈ 281 / 198 / 190 / 669 "
  "wave-by-wave / pooled). The pooled exporter-only specification yields "
  "a negative linear FSTS_c term (β = −0.861, p < .001) but a non-"
  "significant quadratic term (FSTS_c² β = −0.200, p = .660), and the "
  "joint M8 test is not significant (joint F p = .462). We report this "
  "honestly: \"The substantive H1 claim is therefore best read as a non-"
  "monotonic association between participation-and-intensity in "
  "exporting and productivity, not as a strict within-exporter "
  "intensity-curvature claim.\"")
P("• New density-around-turning-point check. Defining a ±5 percentage-"
  "point band around the wave-specific turning point, the share of firms "
  "with FSTS within the band is 0.6 % (6 of 989) in 2009, 1.2 % (11 of "
  "956) in 2015, 0.9 % (9 of 1,013) in 2023, and 1.0 % (29 of 2,958) "
  "pooled. A new auxiliary table (table_density_around_tp.csv) reports "
  "these counts. The §4.5 prose now states: \"the bulk of mass lies at "
  "FSTS = 0; the turning-point estimates are identified primarily through "
  "the contrast between non-exporters and exporters and the right tail "
  "above the TP, not through dense within-band variation.\"")
P("")

# A5 ------------------------------------------------------------------------
H2("A5. Moderation evidence is marginal; need predicted curves and multiple-testing caution")
P("Reviewer comment.", italic=True)
Q('"H4 and the originality claim depend heavily on DAI_z moderation, but '
  'the evidence currently reported is marginal: in 2023 the M8 joint test '
  'is p = .062, the pooled M8 joint test is p = .083… Predicted curves or '
  'marginal-effect plots for low and high DAI_z and TCI_z would clarify '
  'the substantive pattern… clearer multiple-testing caution would also '
  'improve the credibility of the interpretation. If the signal remains '
  'at the current level, H4 may be better positioned as an exploratory '
  'or emerging pattern concentrated in 2023."')
P("Our response.", bold=True)
P("We have implemented all three suggestions:")
P("• New Figure 3 (\"Predicted I–P curves at low / high DAI and TCI\"). "
  "Two panels: 3a shows the FSTS curve at p25 vs p75 of DAI_z; 3b shows "
  "the FSTS curve at p25 vs p75 of TCI_z. Curves are drawn from the "
  "pooled M8 specification with controls and wave + sector fixed effects "
  "fixed at sample modes / means. The figure is embedded inline at the "
  "end of §4.5 (PDF + PNG in submission/figures/).")
P("• Multiple-testing caveat added at the end of §4.5: \"§4.5 reports "
  "nine robustness panels (A–I) across two enriched composites, three "
  "sub-samples, two selection corrections, cross-wave z-tests and a "
  "wave-saturated interaction test, each touching multiple focal terms. "
  "We do not apply a formal multiple-testing correction because the "
  "panels probe different identification concerns rather than testing "
  "the same hypothesis repeatedly, but readers should weight any single "
  "marginal panel result accordingly. Our substantive inferences in §5 "
  "rely on the pattern across panels and the directional consistency of "
  "the focal estimates rather than on the significance of any single "
  "robustness panel.\"")
P("• H4 has been reframed as exploratory and concentrated in 2023: "
  "\"H4 (exploratory). The performance relevance of website-based digital "
  "presence (DAI_z) in Vietnam varies across stages of "
  "internationalisation; we treat the within-wave moderation as "
  "exploratory and concentrate the test on 2023, where the post-NDTP "
  "environment plausibly enables a moderation channel.\" The §4.3 "
  "summary paragraph and the Originality / Value statement have been "
  "edited consistently.")
P("")

# ---------------------------------------------------------------------------
H1("Part B — Detailed Feedback")

H2("B1. Discrepancy in the 2015 and pooled sample sizes (issue #1)")
P("Reviewer comment.", italic=True)
Q('"§§3.1–3.2 and the descriptive table report N = 958 for 2015 and '
  'N = 2,960 pooled, whereas the abstract, Figures 2b/2d, Table LM, and '
  'Table 3 use N = 956 and N = 2,958."')
P("Our response.", bold=True)
P("Real bug. The correct samples under the v5.7 specification are 989 / "
  "956 / 1,013 (pooled 2,958). All occurrences of 958 / 2,960 in the "
  "manuscript prose and the descriptive table have been replaced. The "
  "Table 1 column header now reads \"2015 (N = 956)\" / \"Pooled (N = "
  "2,958)\"; the §3.1 sample-size paragraph, §3.2 missing-code paragraph, "
  "and abstract are all consistent at 989 / 956 / 1,013 / 2,958.")
P("")

H2("B2. IMR vs generalised residual interpretation (issue #2)")
P("Reviewer comment.", italic=True)
Q('"In a probit selection model, the generalised residual coincides with '
  'the IMR for selected observations, so the reported difference in '
  'significance is more plausibly due to the correction being used in a '
  'different sample or specification than to the residual detecting a '
  'distinct selection channel."')
P("Our response.", bold=True)
P("Valid critique. §4.5 Panel E has been rewritten. The key sentence now "
  "reads: \"We do not interpret this gap as evidence of a different "
  "selection channel: in a probit, the generalised residual coincides "
  "with the inverse Mills ratio for selected observations, so the "
  "difference in significance reflects the larger estimation sample "
  "(non-exporters are included in the control-function regression) and "
  "the slightly different specification, not a substantively distinct "
  "selection mechanism. Either way, the focal H1 / H2 / H3 inferences "
  "survive both corrections.\" The earlier wording (\"control-function "
  "signal indicates a degree of selection structure that the standard "
  "IMR does not pick up\") has been removed.")
P("")

H2("B3. Contradictory p-value in Table 3 Panel C (issue #3)")
P("Reviewer comment.", italic=True)
Q('"Panel C gives p = .013 even though it is described as re-estimating '
  'the same DAI_z moderation model on the same N = 1,013 sample. Unless '
  'Panel C is actually a different specification, the reported joint p-'
  'value looks incorrect or mislabeled."')
P("Our response.", bold=True)
P("Real inconsistency. The p = .013 was a residual from an older DAI = "
  "mean(c22b, e6) specification. We have re-run the panel on the "
  "identical 2023 N = 1,013 sample under the v5.7 primary specification "
  "(DAI = c22b only). The reconciled M8 joint test is F(3, 998) ≈ 2.50, "
  "p = .062 — exactly matching the main 2023 M8 result. Both Table 3 "
  "Panel C and the §4.5 prose have been updated to reflect this "
  "reconciled value.")
P("")

H2("B4. Stage-contingent moderation overshoots Panel F evidence (issue #4)")
P("Reviewer comment.", italic=True)
Q('"Panel F finds no statistically significant cross-wave differences in '
  'FSTS_c × DAI_z or FSTS_c² × DAI_z, [so] the claim of a wave-specific '
  'shift in DAI moderation (H4) [is weakened]… Specifically, the 2023 '
  'result is better described as the only wave in which moderation is '
  'detectable within-sample, not as clear evidence of a cross-wave '
  'structural change in the moderation coefficients."')
P("Our response.", bold=True)
P("Valid. §4.3, §4.5, and §5 have been recalibrated. The §4.3 summary "
  "now states: \"H4 receives exploratory support concentrated in the "
  "2023 wave… We do not treat this as a confirmed cross-wave moderation "
  "pattern: the formal pooled wave × focal interaction test (Panel I) "
  "does not detect cross-wave differences in the FSTS × DAI moderation "
  "terms. The substantive reading is that 2023 is the only wave in which "
  "the digital moderation is within-sample detectable, and we frame this "
  "as exploratory rather than as a confirmed hypothesis test.\" The "
  "Originality / Value statement and §5 narrative have been edited "
  "consistently.")
P("")

H2("B5. Contradictory coefficient for the 2015 quadratic (issue #5)")
P("Reviewer comment.", italic=True)
Q('"The 2015 quadratic term is reported inconsistently within §4.1: the '
  'earlier M2 discussion and Table 1 give −2.115 (p = .004), while this '
  'institutional-reading paragraph gives −2.082 (p = .005)."')
P("Our response.", bold=True)
P("Real bug. −2.082 was a leftover value from an old specification. The "
  "institutional-reading paragraph in §4.1 has been corrected to read: "
  "\"the I–P curvature is unusually sharp (FSTS_c² = −2.115, p = .004 in "
  "M2)\". The value now matches the M2 estimate from "
  "coefs_main_models.csv and Table 1.")
P("")

H2("B6. \"Firms added to the exporter cohort\" implies a firm panel (issue #6)")
P("Reviewer comment.", italic=True)
Q('"With repeated cross-sections, the paper cannot identify actual entry '
  'into the exporter cohort, only differences in exporter composition '
  'across waves."')
P("Our response.", bold=True)
P("Valid. The phrase has been replaced. §4.1 now reads: \"…differences "
  "in exporter composition across waves are reflected in productivity "
  "gains coming primarily from scale and from foreign-technology / "
  "standards capability rather than from foundational digital adoption "
  "in this wave.\" §6 (Limitations) restates explicitly that WBES is a "
  "repeated cross-section and within-firm change is not identified.")
P("")

# ---------------------------------------------------------------------------
H1("Summary of new artefacts in v5.7")
P("Robustness panels (Table 3, §4.5):")
P("• Panel C — Common-N reconciled (2023): M8 joint p = .062 (matches main).")
P("• Panel H — Exporter-only sub-sample (FSTS > 0; pooled N = 669): "
  "FSTS_c = −0.861***, FSTS_c² = −0.200 n.s., M8 joint p = .462.")
P("• Panel I — Pooled wave × focal interaction: only DAI_z × wave detectable "
  "(joint p = .016).")
P("• Density-around-TP table: ~1.0 % of pooled firms within ±5 pp band.")
P("Figures:")
P("• Figure 3 — Predicted I–P curves at p25 vs p75 of DAI_z (3a) and "
  "TCI_z (3b), embedded inline at the end of §4.5.")
P("Prose-only revisions: 10 changes covering DAI re-theorise, TCI relabel, "
  "lifecycle recalibration, proxy-obsolescence reading (§5.2.1), H4 "
  "exploratory framing, Panel E rewrite, exporter-cohort wording, multiple-"
  "testing caveat, sample-size reconciliation, and 2015 quadratic fix.")
P("")
P("We thank the Reviewer once again for the thorough engagement with the "
  "manuscript. We believe the revised v5.7 substantively addresses every "
  "comment within the data scope of the WBES Vietnam waves and is ready "
  "for re-evaluation.")
P("")
P("Sincerely,")
P("Do Thuy Huong (Author)")
P("Phan Anh Tu (Corresponding author)")
P("Can Tho University, Vietnam")

doc.save(str(OUT))
print(f"Wrote {OUT}")
