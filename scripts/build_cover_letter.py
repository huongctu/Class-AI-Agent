"""Build cover_letter_ijoem.docx — Emerald-style cover letter for IJoEM.

Run from repo root:
    python3 scripts/build_cover_letter.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE = Path(
    "/root/.claude/uploads/baf64a94-2e13-4dbf-8dcc-22587a547e56/"
    "1f16255e-manuscript_v4_4_1.docx"
)
OUTPUT = Path(__file__).resolve().parents[1] / "cover_letter_ijoem.docx"

STYLE_ID = {
    "Title": "Title", "Subtitle": "Subtitle", "Date": "Date",
    "Heading 1": "Heading1", "Heading 2": "Heading2",
    "First Paragraph": "FirstParagraph",
    "Body Text": "BodyText", "Block Text": "BlockText",
}


def add_para(document, text, style_name):
    para = document.add_paragraph(text)
    pPr = para._p.get_or_add_pPr()
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)
    pStyle.set(qn("w:val"), STYLE_ID[style_name])
    return para


doc = Document(str(TEMPLATE))
body = doc.element.body
for child in list(body):
    if child.tag.endswith("}sectPr"):
        continue
    body.remove(child)
for tbl in list(doc.tables):
    tbl._element.getparent().remove(tbl._element)

add_para(doc, "Cover Letter — International Journal of Emerging Markets",
         "Title")
add_para(doc, "Submission of an original research article", "Subtitle")
add_para(doc, "2026-05-02", "Date")

add_para(doc,
    "Phan Anh Tu (corresponding author)\n"
    "School of Economics (International Business)\n"
    "Can Tho University, Can Tho, Vietnam\n"
    "E-mail: patu@ctu.edu.vn\n"
    "ORCID: https://orcid.org/0000-0003-0667-3137",
    "Body Text")

add_para(doc, "To the Editor-in-Chief,", "Body Text")
add_para(doc, "International Journal of Emerging Markets", "Body Text")
add_para(doc, "Emerald Publishing Limited", "Body Text")

add_para(doc, "Subject of submission", "Heading 2")
add_para(doc,
    'On behalf of my co-author, I am pleased to submit our manuscript, '
    '"Revisiting the Internationalisation–Performance Relationship in an '
    'Emerging Market: The Roles of Technological Capability and Digital '
    'Adoption", for consideration as a research article in the International '
    'Journal of Emerging Markets.',
    "First Paragraph")

add_para(doc, "Fit with the journal", "Heading 2")
add_para(doc,
    "The manuscript is, in our view, a strong fit for IJoEM on three "
    "counts. First, it directly addresses how firms in an emerging market "
    "convert export expansion into productivity outcomes under conditions "
    "of institutional and digital transition — a question central to the "
    "journal's emerging-markets remit. Second, it offers theory refinement "
    "rather than a single-country description: by separating technological "
    "capability from foundational digital adoption with a non-overlapping "
    "primary measurement strategy, the paper recasts the inverted-U "
    "internationalisation–performance debate in a way that is directly "
    "transferable to other transitional emerging economies. Third, the "
    "study uses Vietnam as a context-rich emerging market that has "
    "experienced rapid but uneven institutional and digital change between "
    "2009 and 2023, providing wave-specific evidence on how the productivity "
    "consequences of internationalisation evolve as the institutional "
    "environment matures.",
    "First Paragraph")

add_para(doc, "Empirical contribution", "Heading 2")
add_para(doc,
    "Empirically, the study uses three waves of World Bank Enterprise "
    "Survey microdata for Vietnam (2009, 2015 and 2023) with analytic "
    "samples of 989, 956 and 1,013 firms respectively (pooled N = 2,958). "
    "We estimate ordinary least squares models with HC1 robust standard "
    "errors, quadratic export-intensity terms and interaction "
    "specifications. The main findings are: (i) the inverted-U I–P "
    "relationship is robustly supported, with the Lind–Mehlum test "
    "rejecting monotonicity in all three waves and in the pooled sample "
    "(p ≤ .013) and turning points clustered between 39 and 46 per cent of "
    "direct-export intensity; (ii) technological capability is positive in "
    "all three waves and in the pooled sample and moderates the curvature "
    "in three of four panels; (iii) basic digital adoption is positive in "
    "2009 and 2023 but null in 2015, with the cross-wave shifts "
    "statistically distinguishable in the Paternoster (1998) z-tests; and "
    "(iv) the digital-adoption moderation channel emerges most clearly in "
    "2023 (FSTS_c × DAI_z = −0.912, p = .043). The substantive contribution "
    "of these results is to show that the productivity returns to "
    "internationalisation in an emerging market depend not only on export "
    "intensity but also on how firms combine deeper capability accumulation "
    "with basic digital adoption under transitional institutional "
    "conditions.",
    "First Paragraph")

add_para(doc, "Originality and ethical declarations", "Heading 2")
add_para(doc,
    "This manuscript is original, has not been published previously and is "
    "not under consideration elsewhere. All authors have approved the "
    "submitted version. The authors declare no conflict of interest. The "
    "research received no specific grant from any funding agency in the "
    "public, commercial or not-for-profit sectors. The study uses World "
    "Bank Enterprise Surveys data under the applicable data-access "
    "protocol; accordingly, the raw .dta files are not redistributed, "
    "although replication materials and computational documentation can be "
    "shared to the extent permitted by the data-access terms.",
    "First Paragraph")

add_para(doc, "Author and corresponding author information", "Heading 2")
add_para(doc,
    "Authors:\n"
    "  Do Thuy Huong, College of Economics, Can Tho University, Can Tho, "
    "Vietnam — huongp1323001@gstudent.ctu.edu.vn — ORCID 0000-0002-7711-2487\n"
    "  Phan Anh Tu (corresponding), School of Economics (International "
    "Business), Can Tho University, Can Tho, Vietnam — patu@ctu.edu.vn — "
    "ORCID 0000-0003-0667-3137",
    "Body Text")

add_para(doc, "Suggested classification", "Heading 2")
add_para(doc,
    "Paper type: Research paper. Suggested keywords: internationalisation–"
    "performance; emerging markets; digital adoption; technological "
    "capability; Vietnam; firm productivity. JEL classification: F23, O33, "
    "D22, L25, O53.",
    "First Paragraph")

add_para(doc,
    "We thank you and the anonymous reviewers in advance for the time and "
    "consideration given to our submission. We look forward to your "
    "decision.",
    "Body Text")

add_para(doc, "Sincerely,", "Body Text")
add_para(doc, "Phan Anh Tu", "Body Text")
add_para(doc, "Corresponding author", "Body Text")
add_para(doc, "patu@ctu.edu.vn", "Body Text")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(f"Wrote {OUTPUT}")
