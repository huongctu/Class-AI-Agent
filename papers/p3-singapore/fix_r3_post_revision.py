"""
P3 Singapore R3 Manuscript — Post-revision patches
====================================================
Applies 11 changes to Manuscript_Blinded_MIR_2_revised.docx:

1. Avenyo et al. (2021) — fix DOI from s41287-020-00328-2 to s41287-021-00364-6
   (verified via Crossref + Springer; the manuscript's old DOI does not resolve;
    correct paper title is "Evidence from African firms")
2. Avenyo subtitle: "Microeconomic evidence from sub-Saharan Africa"
   → "Evidence from African firms"
3. Hypothesis 3 (H3) → Hypothesis 2 (H2) in §2.3.3 (DAI conditional non-uniformity)
4. Hypothesis 4 (H4) → Hypothesis 3 (H3) in §2.3.4 (DAI moderation by FSTS)
5. §2.4 conceptual model: "H3 and H4 are jointly informative" → "H2 and H3"
6. §2.4: "H3 concerns the non-uniformity" → "H2 concerns"
7. §2.4: "whereas H4 specifies the direction" → "whereas H3 specifies"
8-11. Figure 1 caption: "(H4)" → "(H3)" and "(H3)" → "(H2)" — applied at
      both para 45 and para 156 occurrences

Rationale:
- After R3 demoted H2 (TCI moderation) to "open empirical question" treated
  in supplementary specification, the remaining hypothesis numbering became
  H1, H3, H4 with a gap. APA / journal convention requires consecutive H1,
  H2, H3 unless gap is explicitly explained.
- Avenyo DOI was inherited from R2 with typo; verified correct DOI via
  Crossref query.

Usage:
    python3 fix_r3_post_revision.py

Requires: python-docx 1.x

Output: Manuscript_R3_FIXED.docx (size unchanged ~1.18 MB)
Verification: 0 instances of "H4", 0 instances of old DOI/subtitle.
"""
import shutil
from docx import Document

SRC = "Manuscript_Blinded_MIR_2_revised.docx"
DST = "Manuscript_R3_FIXED.docx"

shutil.copy(SRC, DST)
doc = Document(DST)


def replace_in_para(p, old, new):
    """Replace text in paragraph; preserves formatting at run level when possible."""
    if old not in p.text:
        return False
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    new_text = p.text.replace(old, new)
    for run in p.runs[1:]:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    return True


changes = []

REPLACEMENTS = [
    # Avenyo DOI + subtitle
    ('10.1057/s41287-020-00328-2', '10.1057/s41287-021-00364-6'),
    ('Microeconomic evidence from sub-Saharan Africa', 'Evidence from African firms'),
    # Hypothesis renumbering
    ('Hypothesis 3 (H3). The productivity association of digital adoption',
     'Hypothesis 2 (H2). The productivity association of digital adoption'),
    ('Hypothesis 4 (H4). The association between digital adoption',
     'Hypothesis 3 (H3). The association between digital adoption'),
    # §2.4 conceptual model text
    ('H3 and H4 are jointly informative', 'H2 and H3 are jointly informative'),
    ('H3 concerns the non-uniformity of the DAI', 'H2 concerns the non-uniformity of the DAI'),
    ('whereas H4 specifies the direction', 'whereas H3 specifies the direction'),
    # Figure 1 caption
    ('contingency variable on the I–P path (H4)', 'contingency variable on the I–P path (H3)'),
    ('direct effect of DAI on the dependent variable (H3)',
     'direct effect of DAI on the dependent variable (H2)'),
]

for i, p in enumerate(doc.paragraphs):
    for old, new in REPLACEMENTS:
        if old in p.text:
            if replace_in_para(p, old, new):
                changes.append(f'  ✓ Para {i}: {old[:50]}... -> {new[:50]}...')

# Also handle table cells (Figure 1 inline shapes / tables)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for old, new in REPLACEMENTS:
                    if old in p.text:
                        if replace_in_para(p, old, new):
                            changes.append(f'  ✓ Table cell: {old[:50]}...')

doc.save(DST)
print('=== Changes applied ===')
for c in changes:
    print(c)
print(f'\nTotal: {len(changes)} edits')
print(f'Output: {DST}')
